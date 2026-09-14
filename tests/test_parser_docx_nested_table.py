r"""DOCX 嵌套表静默不可见性锁定（Round 1927，a 优先级）。

edges33 锁了 cell 内图片不可见、Batch 3 锁了 w:tc 内嵌 w:sdt
递归提取；**cell 内嵌套表**（cell.add_table）grep 零覆盖。探针
R1927 实证（_parse_docx 表序列化只走 body 顶层 w:tbl + cell 文本
来自 cell 自身段落，不含嵌套表段落）：

- **N1 嵌套表不成独立元素**：body 顶层只有一个 w:tbl → 元素
  序 [paragraph, table, paragraph]，无第二个 table 元素
- **N2 嵌套表文本完全不可见**：NESTEDLEFT/NESTEDRIGHT 不出现
  在任何元素 content；零告警（纯静默丢失，与脚注家族同型）
- **N3 外层 cell 呈空**：嵌套表所在的 cell 序列化为空字符串
  （'| BL |  |'），嵌套结构信息也不留痕
- chunk 级联：[seq, isolated_table, seq] 干净隔离

判别式：若表提取改为递归（嵌套表并入 cell 文本或另立元素），
N1/N2 断言翻红。
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser
from app.pipeline import process_single


def _build(path: Path) -> None:
    d = docxlib.Document()
    d.add_paragraph("Intro paragraph.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "TL"
    t.cell(0, 1).text = "TR"
    t.cell(1, 0).text = "BL"
    nested = t.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "NESTEDLEFT"
    nested.cell(0, 1).text = "NESTEDRIGHT"
    d.add_paragraph("Outro paragraph.")
    d.save(str(path))


def _parse(tmp_path: Path):
    p = tmp_path / "nested.docx"
    _build(p)
    return FallbackParser().parse(p, compute_file_hash(p))


def test_nested_table_no_separate_element_text_invisible(tmp_path):
    """N1+N2：嵌套表不成独立 table 元素；NESTEDLEFT/NESTEDRIGHT
    不出现在任何元素；零告警。"""
    d = _parse(tmp_path)
    assert [e.type for e in d.elements] == [
        "paragraph", "table", "paragraph"]
    blob = " ".join(e.content or "" for e in d.elements)
    assert "NESTEDLEFT" not in blob
    assert "NESTEDRIGHT" not in blob
    assert d.warnings == []


def test_outer_cell_with_nested_table_renders_empty(tmp_path):
    """N3：嵌套表所在 cell 序列化为空字符串——外层表格 markdown
    恰 '| TL | TR |\\n| --- | --- |\\n| BL |  |'。"""
    d = _parse(tmp_path)
    assert d.elements[1].content == (
        "| TL | TR |\n| --- | --- |\n| BL |  |")


def test_nested_table_page_chunk_cascade(tmp_path):
    """chunk 级联：[seq, isolated_table, seq]——外层表独立成块且
    其文本恰为表格 markdown；前后正文各自成块。"""
    p = tmp_path / "nested.docx"
    _build(p)
    doc, errors = process_single(p, write_json=False)
    assert errors == []
    assert [c.metadata["strategy"] for c in doc.chunks] == [
        "sequential", "isolated_table", "sequential"]
    table_chunk = doc.chunks[1]
    assert table_chunk.text == doc.elements[1].content
    assert table_chunk.source_element_ids == [
        doc.elements[1].element_id]
    assert doc.chunks[0].text == "Intro paragraph."
    assert doc.chunks[2].text == "Outro paragraph."
