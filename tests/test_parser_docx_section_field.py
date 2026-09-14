r"""DOCX 多节文档 section 字段恒 0 + 分节符空段落实体化锁定（Round 1897，a 优先级）。

edges29 只在单节文档锁过 section: 0（headerReference 挂文末
sectPr）；**多节判别式**零覆盖。探针 R1897 实证（`_parse_docx`
:462 `section_idx = 0` 初始化后循环内**无任何递增**——tag 分派只
有 w:p / w:tbl 两分支，w:sectPr 落空）：

- **section 字段结构性死亡**：真实 2 节文档（python-docx
  `len(sections) == 2` 判别前提），分节后的段落与表格 locator 仍
  全部 section: 0——字段存在但永不反映节结构
- **分节符实体化为空段落**：`add_section` 把 sectPr 挂在新段
  pPr 里，该段无文本 → 成 "(空段落)" 元素，**消耗一个
  paragraph_index**（"after break" 拿 index 2 而非 1）
- **零告警**：节结构完全不可见，无任何 WarningRecord
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _multisec_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("before break")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("after break")
    doc.add_table(rows=1, cols=1)
    doc.tables[0].cell(0, 0).text = "post-section table"
    doc.save(path)


def test_docx_section_field_never_increments(tmp_path):
    """判别式：真实 2 节文档（len(sections)==2），分节后的段落与
    表格 locator 仍全部 section: 0——字段永不递增。"""
    path = tmp_path / "multisec.docx"
    _multisec_docx(path)
    assert len(Document(str(path)).sections) == 2  # 真多节前提
    doc = FallbackParser().parse(path, compute_file_hash(path))
    assert [e.source_locator["section"] for e in doc.elements] == [0, 0, 0, 0]


def test_docx_section_break_renders_as_empty_paragraph(tmp_path):
    """分节符实体化：承载 sectPr 的段无文本 → "(空段落)" 元素，消耗
    一个 paragraph_index（"after break" 是 index 2 而非 1）。"""
    path = tmp_path / "multisec.docx"
    _multisec_docx(path)
    doc = FallbackParser().parse(path, compute_file_hash(path))
    break_para = doc.elements[1]
    assert break_para.content == "(空段落)"
    assert break_para.metadata["empty"] is True
    assert break_para.source_locator["paragraph_index"] == 1
    after = doc.elements[2]
    assert after.content == "after break"
    assert after.source_locator["paragraph_index"] == 2


def test_docx_multi_section_table_locator_and_no_warnings(tmp_path):
    """分节后的表格 locator section 仍 0（table_index 从 0 起）；
    整个多节结构零告警。"""
    path = tmp_path / "multisec.docx"
    _multisec_docx(path)
    doc = FallbackParser().parse(path, compute_file_hash(path))
    tbl = doc.elements[3]
    assert tbl.type == "table"
    assert tbl.source_locator == {"table_index": 0, "section": 0}
    assert [w.code for w in doc.warnings] == []
