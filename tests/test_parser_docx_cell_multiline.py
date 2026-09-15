r"""DOCX 表格 cell 多段落——_Cell.text 以 \\n 连接、md 单元格内嵌换行锁定（Round 1949，a 优先级）。

广扫（cell.add_paragraph / 多段 / second line in cell 全形）
实证既有测试 cell 全部单段（edges4:550 多段是 body 级 para
graph_index）。探针 R1949 实证（python-docx `_Cell.text` =
"\\n".join(p.text)，`_rows_to_markdown` 原样进 md）：

- **M1 cell 两段落**：'line one'+'line two' → md 首格内嵌
  '\\n'：'| line one\\nline two | right |\\n| --- | --- |'——
  **破坏 md 表格行语法但 content 逐字保留**；row/col 计数不
  受影响；零告警
- **M2 cell 段内 w:br**：para.text 同样映射 '\\n' → 与 M1
  完全同形态（两来源经 c.text 收敛）
- **M3 尾随空段落被 strip**：'text' + 空 add_paragraph →
  c.text='text\\n' 但每 cell 先 `.strip()`（fallback_parser.py
  :553）→ **尾随 '\\n' 消失**：'| text | right |'——对照 M1
  内部 '\\n' 原样保留（strip 只削两端）

判别式：若 cell 文本改连接符/全剥换行则 M1/M2/M3 全等断言
翻红；若 strip 取消则 M3 '| text |' 翻红。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _parse(build):
    with tempfile.TemporaryDirectory() as td:
        doc = Document()
        doc.add_paragraph("intro")
        tbl = doc.add_table(rows=1, cols=2)
        build(tbl)
        p = Path(td) / "t.docx"
        doc.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def test_two_paragraph_cell_embeds_newline():
    """M1：cell 两段落 → md 首格 'line one\\nline two'、
    row/col 不变、零告警。"""
    def build(tbl):
        tbl.cell(0, 0).text = "line one"
        tbl.cell(0, 0).add_paragraph("line two")
        tbl.cell(0, 1).text = "right"
    d = _parse(build)
    tbl_e = d.elements[1]
    assert tbl_e.type == "table"
    assert tbl_e.content == "| line one\nline two | right |\n| --- | --- |"
    assert tbl_e.metadata["row_count"] == 1
    assert tbl_e.metadata["col_count"] == 2
    assert d.warnings == []


def test_line_break_cell_converges_same_form():
    """M2：cell 段内 w:br → 与两段落完全同形态（经 c.text 收敛）。"""
    def build(tbl):
        p = tbl.cell(0, 0).paragraphs[0]
        p.add_run("before break")
        p.add_run().add_break()
        p.add_run("after break")
        tbl.cell(0, 1).text = "right"
    d = _parse(build)
    tbl_e = d.elements[1]
    assert tbl_e.content == "| before break\nafter break | right |\n| --- | --- |"
    assert d.warnings == []


def test_trailing_empty_paragraph_stripped():
    """M3：'text' + 空 add_paragraph → c.text='text\\n' 但 cell
    先 strip → 尾随 '\\n' 消失（对照 M1 内部 '\\n' 保留）。"""
    def build(tbl):
        tbl.cell(0, 0).text = "text"
        tbl.cell(0, 0).add_paragraph("")
        tbl.cell(0, 1).text = "right"
    d = _parse(build)
    tbl_e = d.elements[1]
    assert tbl_e.content == "| text | right |\n| --- | --- |"
    assert d.warnings == []
