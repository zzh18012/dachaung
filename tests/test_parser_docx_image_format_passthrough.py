r"""DOCX 位图格式透传——JPEG/TIFF 原始字节直存、ext 由 partname 决定（Round 1946，a 优先级）。

grep 实证 test_parser_docx* / test_pipeline_docx* 全部 PNG——
**JPEG/TIFF 等位图格式零覆盖**。探针 R1946 实证（parser 直存
target.blob、ext 取 partname 后缀，无 PIL 转码/嗅探）：

- **D1 JPEG → ext 'jpg'**：python-docx ImageFormat 约定
  /word/media/image1.jpg（非 'jpeg'）；落盘字节与源 blob
  **逐字节全等**（无转码重编码）
- **D2 TIFF → ext 'tiff'**：partname /word/media/image1.tiff；
  同样逐字节全等
- **D3 PNG 对照 + 宿主空段落**：add_picture 自成段落（无文
  本）→ 中间出现 empty=True 段落元素；图片命名前缀
  para1_00（paragraph_index=1）

判别式：若引入格式嗅探/转码则字节全等断言翻红；若 ext 改取
内容类型映射则 D1 'jpg' 断言翻红；若空宿主段落被跳过则 D3
元素序断言翻红。
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _blob(fmt: str) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), (200, 30, 40)).save(buf, format=fmt)
    return buf.getvalue()


def _parse(tmp_path: Path, fmt: str):
    doc = Document()
    doc.add_paragraph("before")
    doc.add_picture(io.BytesIO(_blob(fmt)))
    p = tmp_path / "d.docx"
    doc.save(p)
    img_dir = tmp_path / "imgs"
    img_dir.mkdir()
    d = FallbackParser(image_output_dir=str(img_dir)).parse(
        p, compute_file_hash(p))
    return d, img_dir


def test_jpeg_passthrough_jpg_ext(tmp_path):
    """D1：JPEG → ext 'jpg'（python-docx 约定非 'jpeg'）、落盘
    字节与源 blob 逐字节全等、零告警。"""
    with tempfile.TemporaryDirectory() as td:
        d, img_dir = _parse(Path(td), "JPEG")
        imgs = [e for e in d.elements if e.type == "image"]
        assert len(imgs) == 1
        e = imgs[0]
        assert e.metadata["ext"] == "jpg"
        assert e.source_locator["target_partname"] == \
            "/word/media/image1.jpg"
        assert next(img_dir.rglob("*.jpg")).read_bytes() == \
            _blob("JPEG")
        assert d.warnings == []


def test_tiff_passthrough_byte_identical(tmp_path):
    """D2：TIFF → ext 'tiff'、字节全等（无转码）。"""
    with tempfile.TemporaryDirectory() as td:
        d, img_dir = _parse(Path(td), "TIFF")
        imgs = [e for e in d.elements if e.type == "image"]
        assert len(imgs) == 1
        e = imgs[0]
        assert e.metadata["ext"] == "tiff"
        assert e.metadata["byte_size"] == len(_blob("TIFF"))
        f = next(img_dir.rglob("*.tiff"))
        assert f.read_bytes() == _blob("TIFF")
        assert d.warnings == []


def test_empty_host_paragraph_and_naming(tmp_path):
    """D3：PNG 对照——图片宿主段落 empty=True 元素在前、命名前缀
    para1_00（paragraph_index=1）、ext 'png'。"""
    with tempfile.TemporaryDirectory() as td:
        d, img_dir = _parse(Path(td), "PNG")
        types = [e.type for e in d.elements]
        assert types == ["paragraph", "paragraph", "image"]
        assert d.elements[1].metadata["empty"] is True
        e = d.elements[2]
        assert e.metadata["ext"] == "png"
        assert e.source_locator["paragraph_index"] == 1
        assert "para1_00" in Path(e.resource_path).name
        assert d.warnings == []
