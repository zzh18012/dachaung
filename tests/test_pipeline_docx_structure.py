r"""pipeline python-docx 标准文档结构端到端（Round 1559）。

新角度：'(空段落)' 占位符进 chunk 文本已在 evaluation
层锁定；本轮锁 **python-docx 标准 API 构建的混合文档**
（真 Heading1 样式 + 空段落 + 正文 + 表格）在 pipeline 层
的完整结构——零覆盖：

- **标题带样式元数据**：heading 元素 level=1、
  style='Heading 1'
- **标题+空段落+正文三元素合并为单 chunk**（文本以
  单空格连接，ids 按序 e0000..e0002）
- **表格 → 独立 markdown chunk** '| cell1 | cell2 |'
  + 分隔行，metadata row_count/col_count/source
- **写盘→validate_only 回环**通过
"""

from __future__ import annotations

import docx as docxlib
from pathlib import Path

from app.pipeline import (
    process_single, validate_only,
)


def _build(tmp_path: Path) -> Path:
    p = tmp_path / "s.docx"
    d = docxlib.Document()
    d.add_heading("Chapter Title", level=1)
    d.add_paragraph("")
    d.add_paragraph("Body text here.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "cell1"
    t.cell(0, 1).text = "cell2"
    d.save(str(p))
    return p


def test_heading_metadata(tmp_path):
    doc, errors = process_single(
        _build(tmp_path),
        write_json=False)
    assert errors == []
    h = doc.elements[0]
    assert h.type == "heading"
    assert h.content == "Chapter Title"
    assert h.metadata["level"] == 1
    assert h.metadata["style"] \
        == "Heading 1"


def test_merged_chunk_and_table_chunk(
        tmp_path):
    doc, errors = process_single(
        _build(tmp_path),
        write_json=False)
    assert errors == []
    prefix = doc.document_id
    (c1, c2) = doc.chunks
    assert c1.text == (
        "Chapter Title (空段落)"
        " Body text here.")
    assert c1.source_element_ids == [
        f"{prefix}::e0000",
        f"{prefix}::e0001",
        f"{prefix}::e0002"]
    assert c2.text == (
        "| cell1 | cell2 |"
        "\n| --- | --- |")
    assert c2.source_element_ids == [
        f"{prefix}::e0003"]
    tab = doc.elements[3]
    assert tab.metadata["row_count"] == 1
    assert tab.metadata["col_count"] == 2
    assert tab.metadata["source"] \
        == "python-docx"


def test_write_validate_roundtrip(
        tmp_path):
    p = _build(tmp_path)
    out = tmp_path / "o.json"
    _, errors = process_single(
        p, out, write_json=True)
    assert errors == []
    ok, msg = validate_only(out)
    assert ok is True
    assert msg == "OK"
