r"""DOCX mc:AlternateContent Choice/Fallback 可见性（Round 1965，a 优先级）。

真实 Word 对 wps 内容必写 mc:AlternateContent（Choice=新版
drawing、Fallback=老版 w:pict）——两分支常含**同一文本**。
edges38 锁过裸 w:pict 文本框不可见、裸 w:sdt 段落不可见，
但 mc: 包装层（Choice/Fallback 双分支语义）零覆盖。风险：
若文本收集按后代 w:t 收集 → 两分支文本双重提取。探针
R1965 实证（fallback 走 body.iterchildren + para.text，
mc: 分支内容**全不可见、无重复、零告警**）：

- **P1 run 级 mc: 两分支各裸 w:t**（机制探针）→ 'before
  after'（两分支全丢；中间双空格保留）
- **P2 run 级真实形态**：Choice=wps 文本框 'CHOICE BOX' /
  Fallback=pict 文本框 'FALLBACK BOX' → 同 P1（两分支全
  丢，无任何 BOX 残留）
- **P3 block 级**：AlternateContent 直接包两 w:p（各含分
  支文本）→ 只剩前后正文段，paragraph_index 顺延 0/1 不
  被包装层消耗

判别式：若文本收集改后代 w:t（.//w:t）则 P1 变 'before
CHOICE TEXT…FALLBACK TEXT… after' 翻红；若实现 Choice
优先可见则 P2 出 'CHOICE BOX' 翻红。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser

_N = (nsdecls("w")
      + ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
      + ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
      + ' xmlns:v="urn:schemas-microsoft-com:vml"')


def _ac(choice_inner: str, fallback_inner: str) -> str:
    return (f"<mc:AlternateContent {_N}>"
            f'<mc:Choice Requires="wps">{choice_inner}</mc:Choice>'
            f"<mc:Fallback>{fallback_inner}</mc:Fallback>"
            f"</mc:AlternateContent>")


def _para_of(text: str) -> str:
    return (f'<w:p {_N}><w:r><w:t xml:space="preserve">{text}'
            f"</w:t></w:r></w:p>")


def _parse(kind: str):
    d = Document()
    if kind == "run_plain":
        p = d.add_paragraph()
        p.add_run("before ")
        p._p.append(parse_xml(
            "<w:r " + _N + ">"
            + _ac("<w:t>CHOICE TEXT</w:t>", "<w:t>FALLBACK TEXT</w:t>")
            + "</w:r>"))
        p.add_run(" after")
    elif kind == "run_textboxes":
        p = d.add_paragraph()
        p.add_run("before ")
        choice = ('<w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
                  '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                  '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
                  "<wps:wsp><wps:txbx><w:txbxContent>"
                  + _para_of("CHOICE BOX")
                  + "</w:txbxContent></wps:txbx></wps:wsp>"
                  "</a:graphicData></a:graphic></wp:anchor></w:drawing>")
        fallback = ('<w:pict><v:shape style="width:100pt;height:20pt">'
                    "<v:textbox><w:txbxContent>"
                    + _para_of("FALLBACK BOX")
                    + "</w:txbxContent></v:textbox></v:shape></w:pict>")
        p._p.append(parse_xml("<w:r " + _N + ">"
                              + _ac(choice, fallback) + "</w:r>"))
        p.add_run(" after")
    else:  # block
        d.add_paragraph("before block")
        d.element.body.append(parse_xml(
            _ac(_para_of("BLOCK CHOICE"), _para_of("BLOCK FALLBACK"))))
        d.add_paragraph("after block")
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "a.docx"
        d.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def test_run_level_plain_text_both_branches_invisible():
    """P1：run 级 mc: 两分支各裸 w:t → 全不可见 → 'before
    after'（中间双空格保留）、零告警。"""
    d = _parse("run_plain")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "before  after"
    assert "CHOICE" not in d.elements[0].content
    assert "FALLBACK" not in d.elements[0].content
    assert d.warnings == []


def test_run_level_dual_textboxes_invisible():
    """P2：真实形态 Choice=wps 文本框 / Fallback=pict 文本框
    → 两分支全丢、无重复提取。"""
    d = _parse("run_textboxes")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "before  after"
    assert "BOX" not in d.elements[0].content
    assert d.warnings == []


def test_block_level_wrapper_skips_paragraphs():
    """P3：block 级 AlternateContent 包两 w:p → 均不可见，
    paragraph_index 顺延 0/1。"""
    d = _parse("block")
    assert [e.content for e in d.elements] == ["before block", "after block"]
    assert [e.source_locator["paragraph_index"] for e in d.elements] == [0, 1]
    assert not any("CHOICE" in (e.content or "") or "FALLBACK" in (e.content or "")
                   for e in d.elements)
    assert d.warnings == []
