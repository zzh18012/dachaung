r"""DOCX 表格行列基数退化与嵌套 w:p 静默丢弃（Round 2029，a 优先级）。

近邻轮全是 PDF 退化族（R2025 操作数栈污染 / R2026 注释吞行与
数字词法 / R2028 文档分诊）；DOCX 侧 R1957（cell 内 hyperlink）、
R1967（gridBefore/gridAfter 列偏移）、R1889（w:sdt 整体丢弃）、
R1899（blip 关系守卫）均未触碰**行列基数退化**：w:tr 无 w:tc、
w:tbl 无 w:tr、w:p 嵌套 w:p、w:p 直接挂 w:tbl（不入 tr）。
探针 R2029 实证（python-docx 1.2.0，_Row.cells 只产 tc_lst、
CT_P.text 只拼 ./w:r|./w:hyperlink 直系子元素）：

- **C1 尾行零 cell**（表头 3 格 + 无格行）→ row.cells=() → 空
  列表行被右补齐 → 尾 body 行 '|  |  |  |'，row_count=2、
  col_count=3、零告警
- **C2 全表仅一行且零 cell** → width=0 仍渲染两行管道空壳
  content='|  |\n|  |'（minLength 1 满足，元素合法落库），
  row_count=1、col_count=0；前置正常表时 table_index=1 正常递增
- **C3 w:tbl 零 w:tr（行空表）** → _rows_to_markdown 返回 "" →
  Element(content="") 触发 models.py 不变量 → **裸 ValueError
  （非 ParserError）穿透 _parse_docx**；pipeline 层
  process_single 兜成 `unexpected_parser_error` + 整篇失败
  （doc=None）——**一张行空表废掉整篇文档**（产品缺陷特征
  锁定，记录于 STATE，不修）
- **C4 w:p 嵌套 w:p**（body 段落直系子含内层 w:p）→ 外层
  text='OUTAFTER'（直系 run 按序拼接）、内层 'IN' 静默丢弃、
  零告警、段落计数仅 1（内层不独立成段）
- **C5 首行零 cell + 次行有内容** → 表头取 norm[0]=补齐空行
  '|  |  |'，真实内容降级为 body 行——表头信息被空行顶掉
- **C6 w:p 直接挂 w:tbl**（不入 w:tr）→ Table.rows 只读
  tr_lst → 'LOST' 整段静默丢弃、表格照常、零告警（与 R1889
  sdt 丢弃同机制不同通道）

判别式：C3 若空 md 被前置守卫拦下或 ValueError 转 ParserError
则翻红；C1/C2/C5 若 _rows_to_markdown 跳过零宽/空行或表头改取
首个非空行则翻红；C4 若段落文本改走 itertext/descendants 则
'IN' 出现翻红；C6 若表格遍历改任何子元素则 'LOST' 出现翻红。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser
from app.pipeline import process_single

_N = nsdecls("w")


def _parse(build):
    """构造 Document → 存盘 → FallbackParser 解析（真实文件路径）。"""
    d = Document()
    build(d)
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "p.docx"
        d.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def _strip_tcs(tr) -> None:
    for tc in list(tr.findall(qn("w:tc"))):
        tr.remove(tc)


def test_cellless_trailing_row_renders_empty_body_row():
    """C1：表头 3 格 + 追加零 cell 的 w:tr → 空 body 行右补齐、
    row_count=2 / col_count=3、零告警。"""

    def build(d: Document) -> None:
        t = d.add_table(rows=1, cols=3)
        for i, h in enumerate(("H1", "H2", "H3")):
            t.rows[0].cells[i].text = h
        t._tbl.append(parse_xml(f"<w:tr {_N}></w:tr>"))

    d = _parse(build)
    assert [e.type for e in d.elements] == ["table"]
    assert d.elements[0].content == (
        "| H1 | H2 | H3 |\n| --- | --- | --- |\n|  |  |  |")
    assert d.elements[0].metadata == {
        "row_count": 2, "col_count": 3, "source": "python-docx"}
    assert d.warnings == []


def test_only_cellless_row_two_pipe_shell():
    """C2：全表仅一行且零 cell → width=0 渲染两行空管道壳
    '|  |\\n|  |'、col_count=0；前置正常表时 table_index=1。"""

    def build(d: Document) -> None:
        t0 = d.add_table(rows=1, cols=2)
        for i, h in enumerate(("A", "B")):
            t0.rows[0].cells[i].text = h
        t1 = d.add_table(rows=1, cols=2)
        _strip_tcs(t1.rows[0]._tr)

    d = _parse(build)
    assert [e.type for e in d.elements] == ["table", "table"]
    shell = d.elements[1]
    assert shell.content == "|  |\n|  |"
    assert shell.metadata == {
        "row_count": 1, "col_count": 0, "source": "python-docx"}
    assert shell.source_locator == {"table_index": 1, "section": 0}
    assert d.warnings == []


def test_rowless_table_kills_whole_document():
    """C3：w:tbl 零 w:tr → 空 md 触发 Element 不变量 → 裸
    ValueError 穿透 parser；pipeline 兜成 unexpected_parser_error
    整篇失败（特征锁定：缺陷记录于 STATE，不修）。"""

    def build(d: Document) -> None:
        d.add_paragraph("keep")
        t = d.add_table(rows=1, cols=2)
        for tr in list(t._tbl.findall(qn("w:tr"))):
            t._tbl.remove(tr)

    d = Document()
    build(d)
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / "in.docx"
        d.save(src)
        with pytest.raises(ValueError, match="必须至少有 content 或 resource_path"):
            FallbackParser().parse(src, compute_file_hash(src))
        doc, errors = process_single(src, td / "out.json", write_json=False)
        assert doc is None
        assert [e.code for e in errors] == ["unexpected_parser_error"]
        assert errors[0].message.startswith("ValueError: element doc-")


def test_nested_wp_inner_text_dropped():
    """C4：body 段落嵌套 w:p → 外层直系 run 按序拼 'OUTAFTER'、
    内层 'IN' 静默丢弃、单段落计数、零告警。"""

    def build(d: Document) -> None:
        p = d.add_paragraph()
        p._p.append(parse_xml(f'<w:r {_N}><w:t>OUT</w:t></w:r>'))
        p._p.append(parse_xml(f'<w:p {_N}><w:r><w:t>IN</w:t></w:r></w:p>'))
        p._p.append(parse_xml(f'<w:r {_N}><w:t>AFTER</w:t></w:r>'))

    d = _parse(build)
    assert [e.type for e in d.elements] == ["paragraph"]
    para = d.elements[0]
    assert para.content == "OUTAFTER"
    assert "IN" not in para.content
    assert para.source_locator == {"paragraph_index": 0, "section": 0}
    assert d.warnings == []


def test_cellless_first_row_becomes_empty_header():
    """C5：首行零 cell + 次行两格内容 → 表头被空补齐行占据
    '|  |  |'、真实内容降级 body 行。"""

    def build(d: Document) -> None:
        t = d.add_table(rows=2, cols=2)
        _strip_tcs(t.rows[0]._tr)
        t.rows[1].cells[0].text = "a"
        t.rows[1].cells[1].text = "b"

    d = _parse(build)
    assert d.elements[0].content == "|  |  |\n| --- | --- |\n| a | b |"
    assert d.elements[0].metadata["row_count"] == 2
    assert d.elements[0].metadata["col_count"] == 2
    assert d.warnings == []


def test_bare_wp_inside_tbl_dropped():
    """C6：w:p 直接挂 w:tbl（不入 w:tr）→ 'LOST' 整段静默丢弃、
    表格照常两列、零告警（tr_lst 之外的子元素不可见）。"""

    def build(d: Document) -> None:
        t = d.add_table(rows=1, cols=2)
        for i, h in enumerate(("A", "B")):
            t.rows[0].cells[i].text = h
        t._tbl.append(parse_xml(f'<w:p {_N}><w:r><w:t>LOST</w:t></w:r></w:p>'))

    d = _parse(build)
    assert [e.type for e in d.elements] == ["table"]
    assert d.elements[0].content == "| A | B |\n| --- | --- |"
    assert all("LOST" != (e.content or "") for e in d.elements)
    assert d.warnings == []
