r"""DOCX 样式名判型残余边界锁定（Round 1931，a 优先级）。

edges10 锁内置样式名（H1-H9 透传、Title=1、Intense Quote 原样）；
**basedOn 继承**与 **heading 前缀异名**零覆盖。_is_heading_style
只匹配样式**名**（"title" / "heading" 前缀 + int 解析，:410-422），
探针 R1931 实证：

- **C1 basedOn 继承不可见**：自定义样式 basedOn Heading 1
  （名 "Section Title"）→ **paragraph**（level 0，样式名保留）——
  用户自定义标题样式静默丢失 heading 语义；chunker 无硬边界 →
  三段融合**单 chunk**（真实世界常见的静默结构丢失首锁）
- **C2 "Heading abc" → heading level=1**：heading 前缀命中但
  int 解析失败 → 回退 1（不是 paragraph）
- **C3 "Heading 10" → heading level=10**：无上限钳制（对照
  edges10 的 H1-H9 全透传）

判别式：若判型改为解析 basedOn 继承链，C1 翻红；若 int 解析
失败回退改为拒判，C2 翻红；若加上限钳制（≤9），C3 翻红。
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib
from docx.enum.style import WD_STYLE_TYPE

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser
from app.pipeline import process_single


def _named(tmp_path: Path, name: str) -> Path:
    d = docxlib.Document()
    d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("Weird named line", style=name)
    p = tmp_path / "named.docx"
    d.save(str(p))
    return p


def test_basedon_heading_inheritance_invisible(tmp_path):
    """C1：自定义样式 basedOn Heading 1（名 "Section Title"）→
    paragraph + level 0 + 样式名保留；无硬边界 → 三段融合单
    chunk——heading 语义静默丢失。"""
    d = docxlib.Document()
    st = d.styles.add_style("Section Title", WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = d.styles["Heading 1"]
    d.add_paragraph("Before body text.")
    d.add_paragraph("Custom styled line", style="Section Title")
    d.add_paragraph("After body text.")
    p = tmp_path / "basedon.docx"
    d.save(str(p))
    doc, errors = process_single(p, write_json=False)
    assert errors == []
    styled = doc.elements[1]
    assert styled.type == "paragraph"
    assert styled.metadata["level"] == 0
    assert styled.metadata["style"] == "Section Title"
    assert len(doc.chunks) == 1
    assert doc.chunks[0].text == \
        "Before body text. Custom styled line After body text."


def test_heading_prefix_unparseable_falls_back_level1(tmp_path):
    """C2："Heading abc" → heading + level=1（int 解析失败回退），
    不是 paragraph。"""
    p = _named(tmp_path, "Heading abc")
    d = FallbackParser().parse(p, compute_file_hash(p))
    e = d.elements[0]
    assert e.type == "heading"
    assert e.metadata["level"] == 1
    assert e.metadata["style"] == "Heading abc"


def test_heading_ten_no_upper_clamp(tmp_path):
    """C3："Heading 10" → heading + level=10（无上限钳制；对照
    edges10 的 H1-H9 透传）。"""
    p = _named(tmp_path, "Heading 10")
    d = FallbackParser().parse(p, compute_file_hash(p))
    e = d.elements[0]
    assert e.type == "heading"
    assert e.metadata["level"] == 10
    assert e.metadata["style"] == "Heading 10"
