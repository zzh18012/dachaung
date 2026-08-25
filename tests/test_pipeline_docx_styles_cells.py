r"""pipeline DOCX 样式层级与 markdown 敌意单元格（Round 1577）。

新角度：R1576 锁表格结构边界——**标题层级全谱、
列表样式降级、单元格特殊文本**零覆盖：

- **标题 level 1..9** → 全部 type='heading'、level=N、
  style='Heading N'
- **List Bullet/Number 样式** → 降级为 type='paragraph'、
  level=0、样式名保留
- **markdown 敌意单元格**：`a|b` 竖线**不转义**、
  换行原样保留、两端空白 strip、空单元格 → ''
- **纯表格文档**（无正文段）→ 单 table 元素 +
  isolated_table chunk
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def test_heading_levels_1_to_9(
        tmp_path):
    d = docxlib.Document()
    for lv in range(1, 10):
        d.add_heading(f"H{lv}", level=lv)
    p = tmp_path / "l.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    assert [(e.type, e.content)
            for e in doc.elements] == [
        ("heading", f"H{lv}")
        for lv in range(1, 10)]
    assert [e.metadata["level"]
            for e in doc.elements] == list(
        range(1, 10))
    assert [e.metadata["style"]
            for e in doc.elements] == [
        f"Heading {lv}"
        for lv in range(1, 10)]


def test_list_styles_downgraded(
        tmp_path):
    d = docxlib.Document()
    d.add_paragraph("Bullet",
                    style="List Bullet")
    d.add_paragraph("Num",
                    style="List Number")
    p = tmp_path / "s.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    assert [(e.type,
             e.metadata["style"],
             e.metadata["level"])
            for e in doc.elements] == [
        ("paragraph", "List Bullet", 0),
        ("paragraph", "List Number", 0)]


def test_markdown_hostile_cells(
        tmp_path):
    d = docxlib.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "a|b"
    t.cell(0, 1).text = ""
    t.cell(1, 0).text = "x\ny"
    t.cell(1, 1).text = "  pad  "
    p = tmp_path / "m.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert el.content == (
        "| a|b |  |"
        "\n| --- | --- |"
        "\n| x\ny | pad |")


def test_table_only_docx(
        tmp_path):
    d = docxlib.Document()
    d.add_table(rows=1, cols=1)\
        .cell(0, 0).text = "solo"
    p = tmp_path / "t.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert el.type == "table"
    assert el.content == (
        "| solo |\n| --- |")
    (c,) = doc.chunks
    assert c.text == el.content
    assert c.metadata["strategy"] == \
        "isolated_table"
