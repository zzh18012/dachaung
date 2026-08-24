r"""app/parsers/fallback_parser.py 边角测试 - 第十八轮（Round 1398）。

新角度（probe 实证）：docx 垂直合并单元格（rowspan，R1382
只锁了水平合并 colspan）：cell(0,0).merge(cell(1,0)) 后文本
在两个行里都出现（'| vmerge | r0c1 |' / '| vmerge | r1c1 |'）
——python-docx 的 merge 是同一 cell 对象被多格引用，extract
时逐格读文本 → 纵向重复。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from app.hash import compute_file_hash
from app.parsers.fallback_parser import \
    FallbackParser


def _vmerge(tmp_path, name="v.docx"):
    d = Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(1, 0))
    t.cell(0, 0).text = "vmerge"
    t.cell(0, 1).text = "r0c1"
    t.cell(1, 1).text = "r1c1"
    p = tmp_path / name
    d.save(str(p))
    return p


def _parse(tmp_path):
    p = _vmerge(tmp_path)
    return FallbackParser().parse(
        p, compute_file_hash(p))


# ---------- 渲染 ----------

def test_vertical_merge_duplicated(tmp_path):
    doc = _parse(tmp_path)
    assert doc.elements[0].content == (
        "| vmerge | r0c1 |\n"
        "| --- | --- |\n"
        "| vmerge | r1c1 |")


def test_single_table_element(tmp_path):
    doc = _parse(tmp_path)
    assert len(doc.elements) == 1
    assert doc.elements[0].type == \
        "table"


def test_table_locator(tmp_path):
    doc = _parse(tmp_path)
    assert doc.elements[
        0].source_locator[
        "table_index"] == 0


def test_no_warnings(tmp_path):
    doc = _parse(tmp_path)
    assert doc.warnings == []


# ---------- schema + 管线 ----------

def test_passes_schema(tmp_path):
    from app.schema import is_valid
    doc = _parse(tmp_path)
    assert is_valid(doc.to_dict())


def test_pipeline_single_chunk(tmp_path):
    from app.pipeline import process_single
    p = _vmerge(tmp_path)
    doc, errors = process_single(
        p, None, parser_name="fallback",
        max_chars=800)
    assert errors == []
    assert len(doc.chunks) == 1
    assert doc.chunks[0].text == (
        "| vmerge | r0c1 |\n"
        "| --- | --- |\n"
        "| vmerge | r1c1 |")
