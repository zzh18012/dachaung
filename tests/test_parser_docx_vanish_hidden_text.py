r"""DOCX w:vanish / w:webHidden 隐藏文本照提（Round 1978，a 优先级）。

模板/表单常把答案、批注藏进 vanish run（Word 显示隐藏文本
需开关）；grep 实证 docx 测试零 vanish/webHidden 匹配（html/
markdown 侧 vanish 是另一语义）。探针 R1978 实证（python-docx
run.text 不看格式属性，隐藏文本与可见**完全同权**）：

- **V1 混排段**：可见 run + vanish run → 'shown SECRET' 拼接
- **V2 整段 vanish** → 'HIDEPARA' 段落照提
- **V3 webHidden run** → 'shown WEBSECRET' 同 vanish
- **V4 表格 cell 内 vanish** → md '| CELLSECRET |' 照提

四态均零告警（无"存在隐藏文本"提示——纯静默提取）。

判别式：若 run.text 按格式过滤隐藏则 V1–V4 文本断言翻红；
若引入 hidden 标记告警则 warnings 断言翻红。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser

_W = nsdecls("w")


def _hidden_run(text: str, prop: str) -> str:
    return (f'<w:r {_W}><w:rPr><w:{prop}/></w:rPr>'
            f'<w:t>{text}</w:t></w:r>')


def _parse(variant: str):
    d = Document()
    if variant == "v1":
        p = d.add_paragraph()
        p.add_run("shown ")
        p._p.append(parse_xml(_hidden_run("SECRET", "vanish")))
    elif variant == "v2":
        p = d.add_paragraph()
        p._p.append(parse_xml(_hidden_run("HIDEPARA", "vanish")))
    elif variant == "v3":
        p = d.add_paragraph()
        p.add_run("shown ")
        p._p.append(parse_xml(_hidden_run("WEBSECRET", "webHidden")))
    else:
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].paragraphs[0]._p.append(
            parse_xml(_hidden_run("CELLSECRET", "vanish")))
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "v.docx"
        d.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def test_vanish_run_concatenated():
    """V1：可见 + vanish run → 'shown SECRET' 拼接、零告警。"""
    d = _parse("v1")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "shown SECRET"
    assert d.warnings == []


def test_whole_paragraph_vanish_extracted():
    """V2：整段 vanish → 'HIDEPARA' 段落照提、零告警。"""
    d = _parse("v2")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "HIDEPARA"
    assert d.warnings == []


def test_webhidden_same_as_vanish():
    """V3：webHidden run → 'shown WEBSECRET' 同 vanish 语义。"""
    d = _parse("v3")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "shown WEBSECRET"
    assert d.warnings == []


def test_vanish_in_table_cell_extracted():
    """V4：表格 cell 内 vanish → md 照提、零告警。"""
    d = _parse("v4")
    assert [e.type for e in d.elements] == ["table"]
    assert d.elements[0].content == "| CELLSECRET |\n| --- |"
    assert d.warnings == []
