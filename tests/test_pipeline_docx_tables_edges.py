r"""pipeline DOCX 表格边界：嵌套/合并与页眉页脚（Round 1576）。

新角度：R1559 锁基础表格——**嵌套表格、合并单元格、
页眉页脚**零覆盖：

- **嵌套表格**（cell.add_table）→ 仅外层表格一个元素，
  内层内容**静默丢弃**（cell(1,0) 为空）
- **合并单元格**（cell.merge）→ 文本在两列**重复出现**，
  row_count/col_count 仍 2×2
- **页眉/页脚** → 完全忽略，仅提取正文
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def test_nested_table_inner_dropped(
        tmp_path):
    d = docxlib.Document()
    outer = d.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "a"
    outer.cell(0, 1).text = "b"
    inner = outer.cell(1, 0).add_table(
        rows=1, cols=2)
    inner.cell(0, 0).text = "x"
    inner.cell(0, 1).text = "y"
    outer.cell(1, 1).text = "c"
    p = tmp_path / "n.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert el.type == "table"
    assert el.content == (
        "| a | b |"
        "\n| --- | --- |"
        "\n|  | c |")
    assert el.metadata["row_count"] == 2
    assert el.metadata["col_count"] == 2
    (c,) = doc.chunks
    assert c.text == el.content
    assert c.metadata["strategy"] == \
        "isolated_table"


def test_merged_cell_duplicated(
        tmp_path):
    d = docxlib.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).text = "merged"
    t.cell(1, 0).text = "p"
    t.cell(1, 1).text = "q"
    p = tmp_path / "m.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert el.content == (
        "| merged | merged |"
        "\n| --- | --- |"
        "\n| p | q |")
    assert el.metadata["row_count"] == 2
    assert el.metadata["col_count"] == 2


def test_header_footer_ignored(
        tmp_path):
    d = docxlib.Document()
    d.sections[0].header.add_paragraph(
        "HDRTEXT")
    d.sections[0].footer.add_paragraph(
        "FTRTEXT")
    d.add_paragraph("Body.")
    p = tmp_path / "h.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    assert [e.content
            for e in doc.elements] == [
        "Body."]
    assert [c.text
            for c in doc.chunks] == [
        "Body."]
