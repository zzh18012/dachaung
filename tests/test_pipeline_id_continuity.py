r"""pipeline id 连续性扫描（Round 1572）。

新角度：chunker 单测锁过单个 id 形态（docX::c0007）；
pipeline 层**全家族 id 连续性**零覆盖：

- element_id 严格 e0000..eNNNN 无缺口
- chunk_id 严格 c0000..cMMMM 无缺口
- 每个 chunk 的 source_element_ids ⊆ element_id
  全集
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib
import pytest

from app.pipeline import process_single
from tests._synthetic_docs import (
    build_minimal_pdf,
)

_FONT = ("<< /Type /Font"
         " /Subtype /Type1"
         " /BaseFont /Helvetica"
         " /Encoding"
         " /WinAnsiEncoding >>")


def _pdf(tmp_path: Path, name: str,
         c: str, with_img=False,
         page2=None) -> Path:
    kids = "[3 0 R]"
    count = 1
    if page2 is not None:
        kids = "[3 0 R 6 0 R]"
        count = 2
    xobj = (" /XObject"
            " << /Im1 8 0 R >>"
            if with_img else "")
    objs = {
        1: "<< /Type /Catalog"
           " /Pages 2 0 R >>",
        2: (f"<< /Type /Pages"
            f" /Kids {kids}"
            f" /Count {count} >>"),
        3: ("<< /Type /Page"
            " /Parent 2 0 R"
            " /MediaBox"
            " [0 0 612 792]"
            " /Resources << /Font"
            f" << /F1 5 0 R >>{xobj} >>"
            " /Contents 4 0 R >>"),
        4: f"<< /Length {len(c)} >>"
           f"\nstream\n{c}\nendstream",
        5: _FONT,
    }
    if page2 is not None:
        objs[6] = ("<< /Type /Page"
                   " /Parent 2 0 R"
                   " /MediaBox"
                   " [0 0 612 792]"
                   " /Resources << /Font"
                   f" << /F1 5 0 R >>{xobj}"
                   " >> /Contents 7 0 R >>")
        objs[7] = (f"<< /Length"
                   f" {len(page2)} >>"
                   f"\nstream\n{page2}\n"
                   f"endstream")
    if with_img:
        px = bytes([255, 0, 0])
        objs[8] = (
            f"<< /Type /XObject"
            f" /Subtype /Image"
            f" /Width 1 /Height 1"
            f" /ColorSpace /DeviceRGB"
            f" /BitsPerComponent 8"
            f" /Length {len(px)} >>"
            f"\nstream\n"
            ).encode() + px \
            + b"\nendstream"
    pdf = b"%PDF-1.4\n"
    for oid in sorted(objs):
        o = objs[oid]
        if isinstance(o, str):
            o = o.encode("latin-1")
        pdf += (f"{oid} 0 obj\n"
                ).encode() + o \
            + b"\nendobj\n"
    pdf += (b"trailer"
            b" << /Root 1 0 R"
            b" /Size "
            + str(len(objs) + 1).encode()
            + b" >>\n%%EOF")
    p = tmp_path / name
    p.write_bytes(pdf)
    return p


def _cases(tmp_path: Path):
    out = {}
    out["plain"] = build_minimal_pdf(
        tmp_path / "p.pdf",
        text="(Hello world)")
    grid = " ".join(
        f"{x} {y} 100 30 re S"
        for x in (72, 172)
        for y in (650, 610))
    out["table"] = _pdf(
        tmp_path, "t.pdf",
        grid + " BT /F1 10 Tf"
        " 80 655 Td (t1) Tj ET"
        " BT /F1 10 Tf 80 615"
        " Td (t2) Tj ET")
    out["image"] = _pdf(
        tmp_path, "i.pdf",
        "q 100 0 0 100 72 692 cm"
        " /Im1 Do Q BT /F1 12"
        " Tf 72 650 Td (BODY)"
        " Tj ET", with_img=True)
    out["multipage"] = _pdf(
        tmp_path, "m.pdf",
        "BT /F1 12 Tf 72 700"
        " Td (P1) Tj ET",
        with_img=True,
        page2="BT /F1 12 Tf 72"
              " 650 Td (P2) Tj ET")
    mixed = tmp_path / "x.docx"
    dd = docxlib.Document()
    dd.add_heading("H", level=1)
    dd.add_paragraph("Body.")
    t = dd.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "a"
    t.cell(0, 1).text = "b"
    dd.save(str(mixed))
    out["mixed-docx"] = mixed
    return out


@pytest.mark.parametrize(
    "name",
    ["plain", "table", "image",
     "multipage", "mixed-docx"])
def test_id_continuity(tmp_path,
                       name):
    p = _cases(tmp_path)[name]
    doc, errors = process_single(
        p, write_json=False,
        max_chars=100)
    assert errors == []
    prefix = doc.document_id
    eids = [e.element_id
            for e in doc.elements]
    assert eids == [
        f"{prefix}::e{i:04d}"
        for i in range(
            len(doc.elements))]
    cids = [c.chunk_id
            for c in doc.chunks]
    assert cids == [
        f"{prefix}::c{i:04d}"
        for i in range(
            len(doc.chunks))]
    eset = set(eids)
    for c in doc.chunks:
        assert set(
            c.source_element_ids) \
            <= eset
