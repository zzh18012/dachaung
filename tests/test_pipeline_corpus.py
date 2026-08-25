r"""pipeline 多文件共享输出目录 + 空 DOCX 警告透传（Round 1560）。

新角度：此前 pipeline 测试全部单文件单输出——**多文件
写入同一输出目录**（images-<sha16> 按内容区分、互不串
扰）与**空 python-docx 文档的警告进 error details** 零
覆盖：

- **4 文件混合语料**（含图 PDF×2 / 纯文本 PDF / DOCX）
  同目录输出：全部校验 OK、document_id 互异、两个
  images-<sha16>/ 各含一个 _p1_00.png（按内容寻址不串）
- **空 python-docx 文档** → no_extracted_elements，警告
  以 details.warnings 透传（docx_no_content）
"""

from __future__ import annotations

import docx as docxlib
from pathlib import Path

from app.pipeline import (
    process_single, validate_only,
)
from tests._synthetic_docs import (
    build_minimal_pdf,
)

_FONT = ("<< /Type /Font"
         " /Subtype /Type1"
         " /BaseFont /Helvetica"
         " /Encoding"
         " /WinAnsiEncoding >>")


def _img_pdf(path: Path, px: bytes,
             label: str) -> None:
    c = (f"q 100 0 0 100 72 692 cm"
         f" /Im1 Do Q"
         f" BT /F1 12 Tf 72 650 Td"
         f" ({label}) Tj ET")
    im = (f"<< /Type /XObject"
          f" /Subtype /Image"
          f" /Width 1 /Height 1"
          f" /ColorSpace /DeviceRGB"
          f" /BitsPerComponent 8"
          f" /Length {len(px)} >>"
          f"\nstream\n"
          ).encode() + px \
        + b"\nendstream"
    objs = {
        1: "<< /Type /Catalog"
           " /Pages 2 0 R >>",
        2: "<< /Type /Pages"
           " /Kids [3 0 R]"
           " /Count 1 >>",
        3: ("<< /Type /Page"
            " /Parent 2 0 R"
            " /MediaBox [0 0 612 792]"
            " /Resources << /Font"
            " << /F1 5 0 R >>"
            " /XObject << /Im1"
            " 7 0 R >> >>"
            " /Contents 4 0 R >>"),
        4: f"<< /Length {len(c)} >>"
           f"\nstream\n{c}\nendstream",
        5: _FONT,
        7: im,
    }
    pdf = b"%PDF-1.4\n"
    for oid in sorted(objs):
        o = objs[oid]
        if isinstance(o, str):
            o = o.encode("latin-1")
        pdf += (f"{oid} 0 obj\n"
                ).encode() + o \
            + b"\nendobj\n"
    pdf += (b"trailer << /Root 1 0 R"
            b" /Size 8 >>\n%%EOF")
    path.write_bytes(pdf)


def test_corpus_shared_outdir(
        tmp_path):
    _img_pdf(tmp_path / "ia.pdf",
             bytes([255, 0, 0]), "REDPIC")
    _img_pdf(tmp_path / "ib.pdf",
             bytes([0, 0, 255]), "BLUEPIC")
    build_minimal_pdf(
        tmp_path / "t.pdf",
        text="(Plain)")
    dp = tmp_path / "m.docx"
    dd = docxlib.Document()
    dd.add_paragraph("Mixed corpus doc.")
    dd.save(str(dp))
    outdir = tmp_path / "all"
    outdir.mkdir()
    ids = set()
    for name in ("ia.pdf", "ib.pdf",
                 "t.pdf", "m.docx"):
        out = outdir / f"{name}.json"
        doc, errors = process_single(
            tmp_path / name, out,
            write_json=True)
        assert errors == []
        ids.add(doc.document_id)
        assert validate_only(out) \
            == (True, "OK")
    assert len(ids) == 4
    dirs = sorted(
        x for x in outdir.iterdir()
        if x.is_dir())
    assert len(dirs) == 2
    for x in dirs:
        assert x.name.startswith(
            "images-")
        files = list(x.iterdir())
        assert [f.name[-10:]
                for f in files] == [
            "_p1_00.png"]
    jsons = sorted(
        x for x in outdir.iterdir()
        if x.suffix == ".json")
    assert [x.name for x in jsons] \
        == ["ia.pdf.json",
            "ib.pdf.json",
            "m.docx.json",
            "t.pdf.json"]


def test_empty_python_docx(
        tmp_path):
    p = tmp_path / "empty.docx"
    docxlib.Document().save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert doc is None
    (e,) = errors
    assert e.code \
        == "no_extracted_elements"
    warns = e.details["warnings"]
    assert [w["code"]
            for w in warns] == [
        "docx_no_content"]
