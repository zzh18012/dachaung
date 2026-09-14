r"""DOCX 行内图中插段落——元素序与文本融合锁定（Round 1933，a 优先级）。

test_pipeline_docx_images 锁资源命名（para0_00 等）；**行内图
与宿主段的元素序、图两侧文本融合形态、locator 关系字段**
零覆盖。探针 R1933 实证：

- **I1 图元素紧跟宿主段之后**：段内行内图 → 元素序
  [paragraph(宿主), image, paragraph(下一段)]——不是图在段前、
  也不是文档末尾集中
- **I2 图切断词无缝愈合**：run "be" + 图 + run "af" → 宿主段
  content 恰 **'beaf'**（图 run 零字符贡献、不插分隔符）；
  有空格时直拼保留 "Before text after text"
- **I3 locator 带宿主关系**：image locator 含宿主段的
  paragraph_index + relationship_id + target_partname；chunk 层
  图零参与（单 sequential 恰宿主段+下一段文本）

判别式：若图元素改排在段前/文档末，I1 序断言翻红；若图 run
被插入占位空格，I2 'beaf' 全等断言翻红。
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def _png() -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
            + chunk(b"IEND", b""))


def _run(tmp_path: Path, split_word: bool):
    d = docxlib.Document()
    p = d.add_paragraph()
    if split_word:
        p.add_run("be")
        p.add_run().add_picture(io.BytesIO(_png()), width=914400)
        p.add_run("af")
    else:
        p.add_run("Before text")
        p.add_run().add_picture(io.BytesIO(_png()), width=914400)
        p.add_run(" after text")
    d.add_paragraph("Next para.")
    p_ = tmp_path / "s.docx"
    d.save(str(p_))
    doc, errors = process_single(p_, write_json=False)
    assert errors == []
    return doc


def test_inline_image_element_follows_host_paragraph(tmp_path):
    """I1：元素序 [paragraph(宿主), image, paragraph(下一段)]——
    图紧跟宿主段之后。"""
    doc = _run(tmp_path, split_word=False)
    assert [e.type for e in doc.elements] == [
        "paragraph", "image", "paragraph"]
    assert doc.elements[0].content == "Before text after text"
    assert doc.elements[2].content == "Next para."


def test_inline_image_splits_word_seamless_fusion(tmp_path):
    """I2：run "be" + 图 + run "af" → 宿主段 content 恰 'beaf'
    （零字符贡献、无分隔符插入）。"""
    doc = _run(tmp_path, split_word=True)
    assert doc.elements[0].content == "beaf"
    assert doc.elements[1].type == "image"


def test_inline_image_locator_and_chunk(tmp_path):
    """I3：image locator 含宿主段 paragraph_index + relationship_id
    + target_partname；图零 chunk（单 sequential 恰两段文本）。"""
    doc = _run(tmp_path, split_word=False)
    img = doc.elements[1]
    loc = img.source_locator
    assert loc["paragraph_index"] == 0
    assert loc["relationship_id"].startswith("rId")
    assert loc["target_partname"].startswith("/word/media/")
    assert len(doc.chunks) == 1
    assert doc.chunks[0].text == "Before text after text Next para."
    assert [e.element_id for e in doc.elements[:1] + doc.elements[2:]] == \
        doc.chunks[0].source_element_ids
