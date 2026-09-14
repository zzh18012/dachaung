r"""BACKLOG §4 变体：w:sdt 包裹内容整体静默丢弃（Round 1889，b 优先级）。

**特征锁定（characterization），不是期望行为规格**。主线
docs/BACKLOG.md §4 记录的缺陷是"批次 14 的 w:sdt 递归覆盖 flow
内容但 w:tc 内 sdt 未纳入"；自跑线分叉早于批次 14（merge-base
2c35244），没有任何 sdt 处理路径——对应**更老的缺陷变体**：
body 级 sdt（段落/表格）与 w:tc 内 sdt 全部静默丢弃，零警告。

根因定位（app/parsers/fallback_parser.py:463）：
`for child in body.iterchildren()` 只按 `w:p`（:465）/ `w:tbl`
（:545）标签分发，`w:sdt` 子元素落入 else 分支 → 无元素、无
警告；表格单元格路径（Table.rows/cells）同样不展开 sdtContent。
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from app.parsers.fallback_parser import _parse_docx

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _sdt(inner_xml: str):
    return parse_xml(
        f'<w:sdt xmlns:w="{_W}">'
        f'<w:sdtPr><w:alias w:val="probe"/></w:sdtPr>'
        f"<w:sdtContent>{inner_xml}</w:sdtContent></w:sdt>"
    )


def _p(text: str) -> str:
    return f'<w:p><w:r><w:t xmlns:w="{_W}">{text}</w:t></w:r></w:p>'


def _make_sdt_docx(path: Path) -> None:
    """before/after 段落 + body 级 sdt 包段落 + body 级 sdt 包表格 +
    普通表格（其 w:tc 内含 sdt 段落，BACKLOG §4 原型）。"""
    d = docx.Document()
    d.add_paragraph("before sdt")
    d.add_paragraph("after sdt")
    body = d.element.body
    after_p = list(body.iterchildren())[1]

    body.insert(list(body).index(after_p), _sdt(_p("inside sdt paragraph")))

    tbl_inner = (
        f'<w:tbl xmlns:w="{_W}">'
        f'<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f'<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
        f'<w:tr><w:tc><w:tcPr/>{_p("sdt cell a")}</w:tc>'
        f'<w:tc><w:tcPr/>{_p("sdt cell b")}</w:tc></w:tr></w:tbl>'
    )
    body.insert(list(body).index(after_p), _sdt(tbl_inner))

    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "plain cell"
    tc = table.cell(0, 1)._tc
    for p in tc.findall(qn("w:p")):
        tc.remove(p)
    tc.append(_sdt(_p("cell inside sdt")))
    d.save(str(path))


def test_backlog_sdt4_body_level_sdt_paragraph_silently_dropped(tmp_path):
    """body 级 sdt 包段落：内容丢弃、零警告、paragraph_index 不为
    其计数（sdt 不存在于解析器视野）。"""
    path = tmp_path / "sdt.docx"
    _make_sdt_docx(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    contents = [e.content for e in elements]
    assert "before sdt" in contents and "after sdt" in contents
    assert "inside sdt paragraph" not in contents
    # 普通段落索引连续 0/1——sdt 段落从未被编号
    paras = [e for e in elements if e.type == "paragraph"]
    assert [p.source_locator["paragraph_index"] for p in paras] == [0, 1]


def test_backlog_sdt4_body_level_sdt_table_silently_dropped(tmp_path):
    """body 级 sdt 包整表：整张表（连单元格文本）丢弃——比主线
    §4 变体（仅 w:tc 边界欠提取）更重的老缺陷。"""
    path = tmp_path / "sdt.docx"
    _make_sdt_docx(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    all_text = "\n".join(e.content for e in elements)
    assert "sdt cell a" not in all_text
    assert "sdt cell b" not in all_text
    # 只剩普通表格那一个 table 元素
    tables = [e for e in elements if e.type == "table"]
    assert len(tables) == 1


def test_backlog_sdt4_cell_level_sdt_text_renders_empty(tmp_path):
    """w:tc 内 sdt（BACKLOG §4 原型）：单元格文本丢弃，markdown
    渲染为空单元格——静默欠提取，零警告。"""
    path = tmp_path / "sdt.docx"
    _make_sdt_docx(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    all_text = "\n".join(e.content for e in elements)
    assert "cell inside sdt" not in all_text
    assert "plain cell" in all_text
    table = next(e for e in elements if e.type == "table")
    lines = table.content.splitlines()
    # 第二列空（sdt 内文本不在），数据行 = '| plain cell |  |'
    assert lines[0] == "| plain cell |  |"
