r"""app/parsers/fallback_parser.py 边角测试 - 第六十一轮（Round 1467）。

新角度（probe 实证）docx 外围 XML 元素可见性（edges49 已锁
fldChar 链缓存结果保留；本轮全是未锁元素）：
- **fldSimple 包裹的内容整体不可见**：TOC 域（含 run 级缓存
  文本 'TOC placeholder'）→ 段落只剩 '(空段落)' 占位
  （empty: True）——与 edges49 的 fldChar separate 段缓存
  **结果保留**形成对比：fldSimple 连缓存都丢
- **w:sym 不可见**（Wingdings 符号 run）：'sym: ' → 'sym:'
- **浮动 wp:anchor 不可见**：w:drawing 无 inline 图片数据 →
  无 image 元素、无告警、段落文本照常
- **core properties 不进 doc.metadata**：title/author 完全
  丢弃（metadata 只有 fallback/image_output_dir 两键）
- **段首 BOM 保留**（﻿ 在 content 里字面保留——docx
  不像 markdown 会杀标题识别，且无 strip）
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.hash import compute_file_hash
from app.parsers.fallback_parser import \
    FallbackParser


def _parse(tmp_path, build, name="p.docx"):
    doc = Document()
    build(doc)
    p = tmp_path / name
    doc.save(str(p))
    return FallbackParser().parse(
        p, compute_file_hash(p))


# ---------- fldSimple ----------

def _add_fldsimple(doc):
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "TOC \\o")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "TOC placeholder"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def test_fldsimple_toc_invisible(
        tmp_path):
    doc = _parse(
        tmp_path, _add_fldsimple)
    assert [(e.content, e.metadata["empty"])
            for e in doc.elements] == [
        ("(空段落)", True),
    ]
    assert "TOC" not in str(
        doc.elements[0].to_dict())


# ---------- w:sym ----------

def _add_sym(doc):
    p = doc.add_paragraph("sym: ")
    r = p.add_run()
    sym = OxmlElement("w:sym")
    sym.set(qn("w:font"), "Wingdings")
    sym.set(qn("w:char"), "F0E0")
    r._r.append(sym)


def test_wsym_invisible(tmp_path):
    doc = _parse(tmp_path, _add_sym)
    assert [e.content
            for e in doc.elements] == ["sym:"]
    assert doc.warnings == []


# ---------- 浮动 anchor ----------

def _add_anchor(doc):
    p = doc.add_paragraph("has drawing")
    r = p.add_run()
    drawing = OxmlElement("w:drawing")
    anchor = OxmlElement("wp:anchor")
    drawing.append(anchor)
    r._r.append(drawing)


def test_floating_anchor_invisible(
        tmp_path):
    doc = _parse(tmp_path, _add_anchor)
    assert [e.content
            for e in doc.elements] == \
        ["has drawing"]
    assert all(
        e.type != "image"
        for e in doc.elements)
    assert doc.warnings == []


# ---------- core properties ----------

def _add_core_props(doc):
    doc.core_properties.title = "My Title"
    doc.core_properties.author = "Author X"
    doc.add_paragraph("body text")


def test_core_props_not_in_metadata(
        tmp_path):
    doc = _parse(tmp_path, _add_core_props)
    assert doc.metadata == {
        "fallback": True,
        "image_output_dir": None,
    }
    assert "My Title" not in str(
        doc.to_dict())


# ---------- BOM ----------

def test_bom_preserved_docx(
        tmp_path):
    doc = _parse(
        tmp_path,
        lambda d: d.add_paragraph(
            "﻿BOM para"))
    assert doc.elements[
        0].content == "﻿BOM para"
    assert doc.elements[
        0].metadata["empty"] is False
