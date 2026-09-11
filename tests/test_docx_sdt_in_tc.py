"""w:tc 内嵌 w:sdt 提取测试（Stage 10 批次 3，BACKLOG §4）。

背景：批次 14 的 w:sdt 递归只覆盖 body 流式层；表格内容走
`_rows_to_markdown` 管线，python-docx 的 _Cell.text 只取 w:tc
直接子级 w:p → 单元格内 sdt 包裹的段落整块静默丢失。

修复：_cell_text 无 sdt 后代走原生 cell.text（逐字节零变化），
有 sdt 走 _iter_cell_paragraphs 文档序递归（保序、不重复提取）。

裁决边界映射（二十四轮授权）：
1. 仅补 w:tc 内内容控件递归路径，不新增语义类别
2. 非 w:tc 路径既有输出零变化（无 sdt 单元格走原生 fast path）
3. 嵌套 w:tbl 不下钻（其内容本就不进 cell.text，非本批范围）
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.schema import validate as validate_udm

docx = pytest.importorskip("docx", reason="python-docx 未安装")
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402


def _sdt_shell():
    return parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:id w:val="1"/></w:sdtPr>'
        f"<w:sdtContent/></w:sdt>"
    )


def _parse(tmp_path: Path, d):
    from app.parsers.fallback_parser import FallbackParser

    p = tmp_path / "synthetic.docx"
    d.save(str(p))
    return FallbackParser().parse(p, source_hash="a" * 64)


def _cell_md(doc) -> str:
    tables = [e for e in doc.to_dict()["elements"] if e["type"] == "table"]
    assert len(tables) == 1
    return tables[0]["content"]


# ---------- 1. 单元格内容整体在 sdt 内 → 不再丢失 ----------

def test_tc_sdt_only_content_extracted(tmp_path: Path):
    d = docx.Document()
    tbl = d.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.text = "S1"
    cell.add_paragraph("S2")
    sdt = _sdt_shell()
    content = sdt.find(qn("w:sdtContent"))
    for p in cell._tc.findall(qn("w:p")):
        content.append(p)  # 移入 sdt
    cell._tc.append(sdt)

    md = _cell_md(_parse(tmp_path, d))
    assert "S1<br>S2" in md  # 段落间 \n 经 linearize_table 渲染为 <br>


# ---------- 2. 混合内容：文档序 + 不重复提取 ----------

def test_tc_mixed_order_no_duplicate(tmp_path: Path):
    d = docx.Document()
    tbl = d.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.text = "P1"
    cell.add_paragraph("S1")
    cell.add_paragraph("S2")
    cell.add_paragraph("P2")
    paras = cell._tc.findall(qn("w:p"))  # [P1, S1, S2, P2]
    sdt = _sdt_shell()
    content = sdt.find(qn("w:sdtContent"))
    for p in paras[1:3]:  # S1、S2 移入 sdt
        content.append(p)
    paras[0].addnext(sdt)  # sdt 放回 P1 之后

    md = _cell_md(_parse(tmp_path, d))
    assert "P1<br>S1<br>S2<br>P2" in md  # 文档序：P1, sdt(S1,S2), P2
    for token in ("P1", "S1", "S2", "P2"):
        assert md.count(token) == 1  # 恰一次，无重复提取


# ---------- 3. sdt 嵌套 sdt（单元格内）→ 递归 ----------

def test_tc_nested_sdt_recursion(tmp_path: Path):
    d = docx.Document()
    tbl = d.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.text = "Deep"
    inner = _sdt_shell()
    for p in cell._tc.findall(qn("w:p")):
        inner.find(qn("w:sdtContent")).append(p)
    outer = _sdt_shell()
    outer.find(qn("w:sdtContent")).append(inner)
    cell._tc.append(outer)

    md = _cell_md(_parse(tmp_path, d))
    assert "Deep" in md
    assert md.count("Deep") == 1


# ---------- 4. 裸 sdt（无 sdtContent）/ 空 sdtContent → 跳过 ----------

def test_tc_bare_sdt_skipped(tmp_path: Path):
    d = docx.Document()
    tbl = d.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.text = "Real"
    bare = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:id w:val="2"/></w:sdtPr></w:sdt>'
    )
    cell._tc.append(bare)
    empty = _sdt_shell()  # sdtContent 存在但为空
    cell._tc.append(empty)

    md = _cell_md(_parse(tmp_path, d))
    assert md.count("Real") == 1  # 裸/空 sdt 不产出也不吞掉真实内容


# ---------- 5. 无 sdt 表格 → 输出与 canonical 契约逐字节一致 ----------

def test_no_sdt_table_byte_identical(tmp_path: Path):
    from app.parsers.table_linearize import linearize_table

    d = docx.Document()
    d.add_paragraph("Intro")
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "H1"
    tbl.cell(0, 1).text = "H2"
    tbl.cell(1, 0).text = "A"
    tbl.cell(1, 1).text = "B|pipe"  # 含竖线的转义路径

    doc = _parse(tmp_path, d)
    md = _cell_md(doc)
    assert md == linearize_table([["H1", "H2"], ["A", "B|pipe"]])
    validate_udm(doc.to_dict())


# ---------- 6. sdt 单元格 + body 段落 → 计数与顺序不受影响 ----------

def test_tc_sdt_with_body_counters(tmp_path: Path):
    d = docx.Document()
    d.add_heading("Title", level=1)
    tbl = d.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.text = "InSdt"
    sdt = _sdt_shell()
    for p in cell._tc.findall(qn("w:p")):
        sdt.find(qn("w:sdtContent")).append(p)
    cell._tc.append(sdt)
    d.add_paragraph("After")

    doc = _parse(tmp_path, d)
    els = doc.to_dict()["elements"]
    assert [e["type"] for e in els] == ["heading", "table", "paragraph"]
    assert "InSdt" in els[1]["content"]
    assert els[1]["source_locator"]["table_index"] == 0
    assert els[2]["source_locator"]["paragraph_index"] == 1
    assert els[1]["metadata"]["row_count"] == 1
    assert els[1]["metadata"]["col_count"] == 1
    validate_udm(doc.to_dict())
