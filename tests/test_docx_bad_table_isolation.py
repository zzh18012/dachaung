"""r55 C08：DOCX 坏表/空表不废整篇文档（解析健壮性）。

- 空表（0 行）跳过留痕：docx_table_empty 结构化 warning（取代批次 5
  裁决④的静默口径，r55 修订），table_index 序不受影响
- 单表提取异常隔离：docx_table_extract_failed warning + 跳过该表，
  其余内容照常产出（不再穿透成 unexpected_parser_error 整篇失败）
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as PyDocx
from docx.oxml.ns import qn

from app.hash import compute_file_hash
from app.parsers import fallback_parser
from app.parsers.fallback_parser import FallbackParser
from app.pipeline import process_single


def _docx_file(tmp_path: Path, build) -> Path:
    d = PyDocx()
    build(d)
    p = tmp_path / "t.docx"
    d.save(str(p))
    return p


def _fill(table, values: list[list[str]]) -> None:
    for i, row in enumerate(values):
        for j, v in enumerate(row):
            table.cell(i, j).text = v


def test_rowless_table_between_tables_leaves_trace(tmp_path: Path):
    """0 行表夹在两张正常表之间：两表照常、空表留痕、索引连续。"""

    def build(d):
        d.add_paragraph("head")
        t0 = d.add_table(rows=1, cols=2)
        _fill(t0, [["A", "B"]])
        t1 = d.add_table(rows=1, cols=2)
        for tr in list(t1._tbl.findall(qn("w:tr"))):
            t1._tbl.remove(tr)
        t2 = d.add_table(rows=1, cols=2)
        _fill(t2, [["C", "D"]])

    doc = FallbackParser().parse(
        _docx_file(tmp_path, build), source_hash="a" * 64)
    tables = [e for e in doc.elements if e.type == "table"]
    assert len(tables) == 2
    assert tables[0].content == "| A | B |\n| --- | --- |"
    assert tables[1].content == "| C | D |\n| --- | --- |"
    # 空表消耗 table_index=1，后续表 index=2（索引语义不变）
    assert tables[0].source_locator["table_index"] == 0
    assert tables[1].source_locator["table_index"] == 2
    assert [(w.code, w.details) for w in doc.warnings] == [
        ("docx_table_empty", {"table_index": 1, "section": 0})]


def test_rowless_table_whole_document_succeeds(tmp_path: Path):
    """pipeline 级回归：0 行表文档 process_single 成功产出（非整篇失败）。"""

    def build(d):
        d.add_paragraph("keep")
        t = d.add_table(rows=1, cols=2)
        for tr in list(t._tbl.findall(qn("w:tr"))):
            t._tbl.remove(tr)

    src = _docx_file(tmp_path, build)
    doc, errors = process_single(src, tmp_path / "out.json", write_json=False)
    assert errors == []
    assert doc is not None
    assert [e.content for e in doc.elements if e.type == "paragraph"] == ["keep"]
    assert [w.code for w in doc.warnings] == ["docx_table_empty"]


def test_table_extract_exception_isolated(tmp_path: Path, monkeypatch):
    """单表提取抛异常：该表跳过 + docx_table_extract_failed 告警，
    其余内容照常（不再穿透成整篇 unexpected_parser_error）。"""

    def build(d):
        d.add_paragraph("before")
        t0 = d.add_table(rows=1, cols=2)
        _fill(t0, [["A", "B"]])
        d.add_paragraph("after")

    src = _docx_file(tmp_path, build)

    real = fallback_parser._rows_to_markdown

    def _boom(rows):
        raise RuntimeError("synthetic table extraction failure")

    monkeypatch.setattr(fallback_parser, "_rows_to_markdown", _boom)
    try:
        doc = FallbackParser().parse(src, compute_file_hash(src))
    finally:
        monkeypatch.setattr(fallback_parser, "_rows_to_markdown", real)
    assert [e.content for e in doc.elements if e.type == "paragraph"] == [
        "before", "after"]
    assert [e.type for e in doc.elements if e.type == "table"] == []
    assert len(doc.warnings) == 1
    w = doc.warnings[0]
    assert w.code == "docx_table_extract_failed"
    assert w.details["exception_type"] == "RuntimeError"
    assert w.details["table_index"] == 0


def test_normal_table_still_no_warnings(tmp_path: Path):
    """回归护栏：正常表格文档零告警（C08 不误伤正常路径）。"""

    def build(d):
        t = d.add_table(rows=2, cols=2)
        _fill(t, [["h1", "h2"], ["a", "b"]])

    doc = FallbackParser().parse(
        _docx_file(tmp_path, build), source_hash="a" * 64)
    assert doc.warnings == []
    assert len(doc.elements) == 1
