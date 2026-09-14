r"""DOCX 表格单元格内嵌图片静默丢失锁定（Round 1948，a 优先级）。

grep 实证无任何 DOCX 测试把图片放进表格 cell（PDF 侧
grid_cell_image 已锁）。parser 的 w:tbl 分支只读 `c.text`
（fallback_parser.py :552-556）——cell 段落不走
`_extract_inline_image_rids`。探针 R1948 实证：

- **T1 cell 图片静默丢失**：1x2 表格 cell(0,1) 内嵌 PNG →
  **零 image 元素、零告警、零 PNG 落盘**（结构静默丢失，与
  edges69 numPr/outlineLvl 同族：Word 语义有、提取面没有）
- **表格 md 空 cell**：'| text cell |  |'——图片占位不进
  单元格文本
- **T2 body 图片不受影响**：同文档 body 段落图片照常成元素
  （恰 1 个、para2_00、PNG 落盘）——丢失只发生在 tbl 分支

判别式：若 tbl 分支补扫 cell 段落 rId 则 T1/T2 元素数断言翻
红；若引入丢失告警则零告警断言翻红；md 全等断言锁 cell 文
本形态。
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _blob() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (2, 2), (10, 200, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _parse(tmp_path: Path, with_body_pic: bool):
    doc = Document()
    doc.add_paragraph("intro")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "text cell"
    run_ = tbl.cell(0, 1).paragraphs[0].add_run()
    run_.add_picture(io.BytesIO(_blob()))
    if with_body_pic:
        doc.add_paragraph("after")
        doc.add_picture(io.BytesIO(_blob()))
    p = tmp_path / "t.docx"
    doc.save(p)
    img_dir = tmp_path / "imgs"
    img_dir.mkdir()
    return FallbackParser(image_output_dir=str(img_dir)).parse(
        p, compute_file_hash(p)), img_dir


def test_cell_image_silently_dropped(tmp_path):
    """T1：cell 内图片 → 零 image 元素、零 PNG、零告警。"""
    with tempfile.TemporaryDirectory() as td:
        d, img_dir = _parse(Path(td), with_body_pic=False)
        assert [e.type for e in d.elements] == ["paragraph", "table"]
        assert not any(e.type == "image" for e in d.elements)
        assert list(img_dir.rglob("*.png")) == []
        assert d.warnings == []


def test_table_md_empty_image_cell(tmp_path):
    """表格 md：图片 cell 呈空 '| text cell |  |'，row/col 1/2。"""
    with tempfile.TemporaryDirectory() as td:
        d, _ = _parse(Path(td), with_body_pic=False)
        tbl = d.elements[1]
        assert tbl.content == "| text cell |  |\n| --- | --- |"
        assert tbl.metadata["row_count"] == 1
        assert tbl.metadata["col_count"] == 2


def test_body_image_survives_alongside(tmp_path):
    """T2：同文档 body 图片照常提取——恰 1 image、para2_00、
    PNG 落盘、零告警。"""
    with tempfile.TemporaryDirectory() as td:
        d, img_dir = _parse(Path(td), with_body_pic=True)
        types = [e.type for e in d.elements]
        assert types == ["paragraph", "table", "paragraph",
                         "paragraph", "image"]
        imgs = [e for e in d.elements if e.type == "image"]
        assert len(imgs) == 1
        assert imgs[0].source_locator["paragraph_index"] == 2
        pngs = list(img_dir.rglob("*.png"))
        assert len(pngs) == 1
        assert "para2_00" in pngs[0].name
        assert d.warnings == []
