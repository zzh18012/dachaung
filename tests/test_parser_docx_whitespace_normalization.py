r"""DOCX 纯空白段 ≡ 空串归一锁定（Round 1928，a 优先级）。

R1925/1926 锁空串 heading；R1585 锁分页符段（w:br 无字符贡献）。
**纯空白文本**（空格/制表符）经 _parse_docx :468
`(para.text or "").strip()` 在 intake 即归一为空串——下游全链
与空段同型（探针 R1928 实证，推翻"content 保留原始空白"的
代码直觉）：

- **W1 空格段**："   " Normal 段 → content=="(空段落)"（非
  原始 "   "）、empty=True
- **W2 占位符入 chunk**：两正文间空白段 → 单 chunk
  'Before body text. (空段落) After body text.'，三个元素
  全进 source_ids（占位符非空白文本，不触发 chunker 跳过）
- **W3 空白 heading**："   " Heading 1 段 → type=heading、
  content=="(空段落)"、level=1、empty=True；两正文间**照常
  劈 2 chunk**（与 R1925 空 heading 同型——判据是 strip 后
  非空的占位符文本，元素类型照样触发硬边界）
- **W4 制表符变体**："\t" 段 → 同占位符（strip 吞掉制表符）

判别式：若 :468 的 strip 移除（或空白判空改为原样保留），
W1/W4 的 content 断言翻红；若占位符被视为空白被 chunker 跳过，
W2 融合断言翻红。
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser
from app.pipeline import process_single


def _build(path: Path, parts: list[tuple[str, str]]) -> None:
    d = docxlib.Document()
    for kind, text in parts:
        if kind == "p":
            para = d.add_paragraph(text)
        else:
            para = d.add_paragraph(text, style="Heading 1")
        assert para.text == text
    d.save(str(path))


def test_space_paragraph_normalized_to_placeholder(tmp_path):
    """W1：纯空格段 → content=="(空段落)"、empty=True——strip 在
    intake 归一，不保留原始空白。"""
    p = tmp_path / "ws.docx"
    _build(p, [("p", "Before body text."), ("p", "   "),
               ("p", "After body text.")])
    d = FallbackParser().parse(p, compute_file_hash(p))
    ws = d.elements[1]
    assert ws.type == "paragraph"
    assert ws.content == "(空段落)"
    assert ws.metadata["empty"] is True


def test_whitespace_placeholder_fuses_all_three_elements(tmp_path):
    """W2：两正文间空白段 → 单 chunk，占位符进 chunk 文本、三元素
    全进 source_ids。"""
    p = tmp_path / "ws2.docx"
    _build(p, [("p", "Before body text."), ("p", "   "),
               ("p", "After body text.")])
    doc, errors = process_single(p, write_json=False)
    assert errors == []
    assert len(doc.chunks) == 1
    assert doc.chunks[0].text == \
        "Before body text. (空段落) After body text."
    assert doc.chunks[0].source_element_ids == [
        e.element_id for e in doc.elements]


def test_whitespace_heading_splits_like_empty_heading(tmp_path):
    """W3：空白 Heading 1 → heading + 占位符 + level=1；照常劈
    2 chunk（与 R1925 空 heading 同型）。"""
    p = tmp_path / "wsh.docx"
    _build(p, [("p", "Before body text."), ("h", "   "),
               ("p", "After body text.")])
    doc, errors = process_single(p, write_json=False)
    assert errors == []
    h = doc.elements[1]
    assert h.type == "heading"
    assert h.content == "(空段落)"
    assert h.metadata["level"] == 1
    assert h.metadata["empty"] is True
    assert [c.text for c in doc.chunks] == [
        "Before body text.", "(空段落) After body text."]


def test_tab_paragraph_normalized_to_placeholder(tmp_path):
    """W4：制表符段 → 同占位符（strip 吞掉 \\t）。"""
    p = tmp_path / "tab.docx"
    _build(p, [("p", "\t")])
    d = FallbackParser().parse(p, compute_file_hash(p))
    assert len(d.elements) == 1
    assert d.elements[0].content == "(空段落)"
    assert d.elements[0].metadata["empty"] is True
