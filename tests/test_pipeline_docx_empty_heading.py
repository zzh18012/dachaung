r"""DOCX 空 heading 级联锁定（Round 1925，a 优先级）。

R1585 锁了**非空** Title/Quote 样式与 Normal 空段（分页符）；
空样式 heading（style 判型 × 空文本占位 × chunk 硬边界）零覆盖。
探针 R1925 实证（_parse_docx 样式判型不看文本是否为空）：

- **样式判型优先于空文本**：空 Heading 1 段 → 元素 type=heading、
  content="(空段落)"、metadata level=1 + empty=True——样式赢了
  判型，空赢了占位文本
- **硬边界按元素类型触发**：两正文段之间的空 heading 仍把流
  劈成 2 个 sequential chunk——零真实文本照样封口前块；占位符
  "(空段落)" 进入第二个 chunk 文本并与其后正文融合
- **Title 样式同型**：空 Title 段 → heading level=1，同样的
  二劈 chunk 序

判别式：若样式判型改为要求非空文本（空样式段降级 paragraph），
判型/劈块测试翻红；若 chunker 过滤空占位文本，融合断言翻红。
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def _build(path: Path, style: str) -> None:
    d = docxlib.Document()
    d.add_paragraph("Before body text.")
    h = d.add_paragraph("", style=style)
    assert h.text == ""
    d.add_paragraph("After body text.")
    d.save(str(path))


def _run(tmp_path: Path, style: str):
    p = tmp_path / "s.docx"
    _build(p, style)
    doc, errors = process_single(p, write_json=False)
    assert doc is not None and errors == []
    return doc


def test_empty_heading1_element_shape(tmp_path):
    """空 Heading 1 段 → type=heading、content="(空段落)"、
    level=1、empty=True（样式判型不看文本空否）。"""
    doc = _run(tmp_path, "Heading 1")
    assert [(e.type, e.content) for e in doc.elements] == [
        ("paragraph", "Before body text."),
        ("heading", "(空段落)"),
        ("paragraph", "After body text."),
    ]
    h = doc.elements[1]
    assert h.metadata["level"] == 1
    assert h.metadata["empty"] is True
    assert h.metadata["style"] == "Heading 1"
    assert doc.elements[0].metadata["empty"] is False


def test_empty_heading_hard_boundary_splits_chunks(tmp_path):
    """空 heading 零真实文本仍触发硬边界：2 个 sequential chunk，
    前块恰 "Before body text."；后块以占位符开头并与后文融合。"""
    doc = _run(tmp_path, "Heading 1")
    assert [c.metadata["strategy"] for c in doc.chunks] == [
        "sequential", "sequential"]
    first, second = doc.chunks
    assert first.text == "Before body text."
    assert first.source_element_ids == [
        doc.elements[0].element_id]
    assert second.text == "(空段落) After body text."
    assert second.source_element_ids == [
        doc.elements[1].element_id,
        doc.elements[2].element_id]


def test_title_style_empty_heading_same_cascade(tmp_path):
    """空 Title 段与空 Heading 1 同型：heading level=1 + 同样的
    二劈 chunk 序（Title 也是 level 1 标题家族）。"""
    doc = _run(tmp_path, "Title")
    h = doc.elements[1]
    assert h.type == "heading"
    assert h.metadata["level"] == 1
    assert h.metadata["empty"] is True
    assert h.metadata["style"] == "Title"
    assert [c.text for c in doc.chunks] == [
        "Before body text.", "(空段落) After body text."]
