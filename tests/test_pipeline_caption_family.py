r"""pipeline 题注 caption 家族：正则变体与隔离（Round 1593）。

新角度：R1591/1592 锁切分边界——**caption 分类
正则与 isolated_caption 策略**在合成文档层零覆盖：

- **PDF 正变体**：Table 1:/Figure 2/Fig. 3./TABLE 4
  （大小写不敏感）→ caption + isolated_caption；
  **负变体**：'Table of contents'（无数字）、
  'Tab 6'（前缀不符）→ heading + sequential
- **DOCX 中文**：表 1:/图 2、/表 3 → caption；
  **全角冒号紧贴数字** '表 1：x' → 分隔符类不含
  '：' → 仍是 paragraph（quirk）
- caption 前后段落不与之合并（隔离成独立 chunk）
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single

_FONT = ("<< /Type /Font"
         " /Subtype /Type1"
         " /BaseFont /Helvetica"
         " /Encoding"
         " /WinAnsiEncoding >>")


def _pdf(tmp_path: Path, name: str,
         c: str) -> Path:
    objs = {
        1: "<< /Type /Catalog"
           " /Pages 2 0 R >>",
        2: "<< /Type /Pages"
           " /Kids [3 0 R]"
           " /Count 1 >>",
        3: ("<< /Type /Page"
            " /Parent 2 0 R"
            " /MediaBox"
            " [0 0 612 792]"
            " /Resources << /Font"
            " << /F1 5 0 R >> >>"
            " /Contents 4 0 R >>"),
        4: f"<< /Length {len(c)} >>"
           f"\nstream\n{c}\nendstream",
        5: _FONT,
    }
    pdf = b"%PDF-1.4\n"
    for oid in sorted(objs):
        o = objs[oid]
        if isinstance(o, str):
            o = o.encode("latin-1")
        pdf += (f"{oid} 0 obj\n"
                ).encode() + o \
            + b"\nendobj\n"
    pdf += (b"trailer"
            b" << /Root 1 0 R"
            b" /Size 6 >>\n%%EOF")
    p = tmp_path / name
    p.write_bytes(pdf)
    return p


def test_pdf_caption_variants(
        tmp_path):
    lines = [
        "Table 1: Results of experiment",
        "Figure 2 shows the trend",
        "Fig. 3. An overview",
        "TABLE 4 Overview",
        "Table of contents goes here now",
        "Tab 6: not a caption marker",
    ]
    parts = []
    y = 750
    for t in lines:
        parts.append(
            f"BT /F1 10 Tf 72 {y}"
            f" Td ({t}) Tj ET")
        y -= 100
    p = _pdf(tmp_path, "c.pdf",
             " ".join(parts))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    types = [e.type
             for e in doc.elements]
    assert types == [
        "caption", "caption",
        "caption", "caption",
        "heading", "heading"]
    got = [(c.metadata["strategy"])
           for c in doc.chunks]
    assert got == [
        "isolated_caption"] * 4 \
        + ["sequential"] * 2


def test_docx_chinese_captions(
        tmp_path):
    d = docxlib.Document()
    d.add_paragraph("表 1: ASCII冒号")
    d.add_paragraph("图 2、顿号")
    d.add_paragraph("表 3 全角：冒号")
    d.add_paragraph("表 1：紧贴全角冒号")
    p = tmp_path / "c.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    got = [(e.type, e.content)
           for e in doc.elements]
    assert got == [
        ("caption", "表 1: ASCII冒号"),
        ("caption", "图 2、顿号"),
        ("caption", "表 3 全角：冒号"),
        ("paragraph",
         "表 1：紧贴全角冒号")]
    caps = [c for c in doc.chunks
            if c.metadata["strategy"]
            == "isolated_caption"]
    assert len(caps) == 3


def test_caption_isolated_between(
        tmp_path):
    lines = [
        "Before caption paragraph text",
        "Table 9: the caption line",
        "After caption paragraph text",
    ]
    parts = []
    y = 750
    for t in lines:
        parts.append(
            f"BT /F1 10 Tf 72 {y}"
            f" Td ({t}) Tj ET")
        y -= 100
    p = _pdf(tmp_path, "s.pdf",
             " ".join(parts))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    got = [(c.text,
            c.metadata["strategy"])
           for c in doc.chunks]
    assert got[1] == (
        "Table 9: the caption line",
        "isolated_caption")
    assert got[0][1] == "sequential"
    assert got[2][1] == "sequential"
    assert all(
        len(c.source_element_ids) == 1
        for c in doc.chunks)
