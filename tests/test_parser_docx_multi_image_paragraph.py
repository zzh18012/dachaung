r"""DOCX 同段多行内图——元素序/命名/同字节复用锁定（Round 1938）。

R1933 锁单图元素序 [para, image, para]；test_pipeline_docx_images
锁单图 para0_00 命名。**同段两图**与**同图字节复用**零覆盖。
探针 R1938 实证：

- **P1 两图都跟宿主段后**：同段 "L " + 图1 + 图2 + " R" →
  元素序 [paragraph('L  R'), image, image, paragraph]——按
  插入序连排宿主段后（不是段前、不是文档末集中）
- **P2 段内顺序命名**：两异图 → image_{sha}_para0_00.png /
  _para0_01.png，rels rId9/rId10，两个 PNG 落盘
- **P3 同字节复用共享 rel**：同 PNG 字节插两次 → **两元素共享
  rId9 + 同 target_partname /word/media/image1.png**（python-docx
  部件去重），但解析器仍各写一份 PNG（磁盘字节重复、
  para0_00/01 各自计数）

判别式：若图元素改排文档末集中则 P1 序断言翻红；若命名按
rel 去重（复用共号）则 P3 文件对断言翻红。
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def _png(rgb: bytes = b"\x00\xff\x00") -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(rgb + b"\x00"))
            + chunk(b"IEND", b""))


def _run(tmp_path: Path, reuse: bool):
    p = tmp_path / "s.docx"
    d = docxlib.Document()
    para = d.add_paragraph()
    para.add_run("L ")
    para.add_run().add_picture(io.BytesIO(_png()), width=914400)
    if reuse:
        para.add_run().add_picture(io.BytesIO(_png()))
    else:
        para.add_run().add_picture(io.BytesIO(_png(b"\x00\x00\xff")))
    para.add_run(" R")
    d.add_paragraph("Next.")
    d.save(str(p))
    doc, errors = process_single(p, output_path=tmp_path / "o" / "s.json")
    assert errors == []
    return doc, tmp_path / "o"


def test_two_images_follow_host_in_insertion_order(tmp_path):
    """P1：同段两图 → [para('L  R'), image, image, para]，按插入
    序连排宿主段后。"""
    doc, _ = _run(tmp_path, reuse=False)
    assert [e.type for e in doc.elements] == [
        "paragraph", "image", "image", "paragraph"]
    assert doc.elements[0].content == "L  R"
    assert doc.elements[3].content == "Next."


def test_sequential_naming_within_paragraph(tmp_path):
    """P2：两异图 → _para0_00 / _para0_01 + 两 PNG 落盘。"""
    doc, out = _run(tmp_path, reuse=False)
    imgs = [e for e in doc.elements if e.type == "image"]
    assert imgs[0].resource_path.endswith("para0_00.png")
    assert imgs[1].resource_path.endswith("para0_01.png")
    assert [e.source_locator["relationship_id"] for e in imgs] == [
        "rId9", "rId10"]
    names = sorted(f.name for f in out.rglob("*.png"))
    assert [n.endswith("para0_00.png") for n in names] == [True, False]
    assert len(names) == 2
    assert all(e.metadata["extracted_to_disk"] for e in imgs)


def test_same_bytes_reuse_shared_rel_two_elements(tmp_path):
    """P3：同字节复用 → 两元素共享 rId9 + 同 partname，但各写
    一份 PNG（磁盘字节重复、各自计数）。"""
    doc, out = _run(tmp_path, reuse=True)
    imgs = [e for e in doc.elements if e.type == "image"]
    assert len(imgs) == 2
    assert [e.source_locator["relationship_id"] for e in imgs] == [
        "rId9", "rId9"]
    assert [e.source_locator["target_partname"] for e in imgs] == [
        "/word/media/image1.png"] * 2
    assert len(list(out.rglob("*.png"))) == 2
    assert imgs[0].element_id != imgs[1].element_id
