r"""pipeline 元素置信度分类学（Round 1594）。

新角度：R1593 锁 caption——**confidence 字段的
parser 级取值**在合成文档层零覆盖：

- **PDF 混合页**（文本 + 表格 + 图片）→
  heading 0.85 / table 0.7 / image 0.6（页内序
  text→table→image，与绘制顺序无关）
- **DOCX 混合**（段落 + 表格 + 图片）→ 全部
  0.95；JSON 往返保留数值
"""

from __future__ import annotations

import json

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single

_FONT = ("<< /Type /Font"
         " /Subtype /Type1"
         " /BaseFont /Helvetica"
         " /Encoding"
         " /WinAnsiEncoding >>")


def _mixed_pdf(tmp_path: Path) -> Path:
    grid = " ".join(
        f"{x} {y} 100 30 re S"
        for x in (72, 172)
        for y in (650, 620))
    c = ("q 100 0 0 100 300 650"
         " cm /Im1 Do Q " + grid
         + " BT /F1 10 Tf 80 625"
         " Td (t1) Tj ET")
    px = bytes([255, 0, 0])
    img = ((f"<< /Type /XObject"
            f" /Subtype /Image"
            f" /Width 1 /Height 1"
            f" /ColorSpace /DeviceRGB"
            f" /BitsPerComponent 8"
            f" /Length {len(px)} >>"
            f"\nstream\n"
            ).encode()
           + px + b"\nendstream")
    objs = {
        1: "<< /Type /Catalog"
           " /Pages 2 0 R >>",
        2: "<< /Type /Pages"
           " /Kids [3 0 R]"
           " /Count 1 >>",
        3: ("<< /Type /Page"
            " /Parent 2 0 R"
            " /MediaBox"
            " [0 0 612 792]"
            " /Resources << /Font"
            " << /F1 5 0 R >>"
            " /XObject"
            " << /Im1 8 0 R >> >>"
            " /Contents 4 0 R >>"),
        4: f"<< /Length {len(c)} >>"
           f"\nstream\n{c}\nendstream",
        5: _FONT,
        8: img,
    }
    pdf = b"%PDF-1.4\n"
    for oid in sorted(objs):
        o = objs[oid]
        if isinstance(o, str):
            o = o.encode("latin-1")
        pdf += (f"{oid} 0 obj\n"
                ).encode() + o \
            + b"\nendobj\n"
    pdf += (b"trailer"
            b" << /Root 1 0 R"
            b" /Size 9 >>\n%%EOF")
    p = tmp_path / "m.pdf"
    p.write_bytes(pdf)
    return p


def test_pdf_confidence_taxonomy(
        tmp_path):
    p = _mixed_pdf(tmp_path)
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    got = [(e.type, e.confidence)
           for e in doc.elements]
    assert got == [
        ("heading", 0.85),
        ("table", 0.7),
        ("image", 0.6)]


def test_docx_confidence_uniform(
        tmp_path):
    import io
    import struct
    import zlib

    def chunk(t, data):
        return (struct.pack(">I",
                            len(data))
                + t + data
                + struct.pack(
                    ">I",
                    zlib.crc32(t + data)
                    & 0xffffffff))

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR",
                   struct.pack(
                       ">IIBBBBB",
                       1, 1, 8, 2,
                       0, 0, 0))
           + chunk(b"IDAT",
                   zlib.compress(
                       b"\x00\xff\x00\x00"))
           + chunk(b"IEND", b""))

    d = docxlib.Document()
    d.add_paragraph("Para.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "c"
    d.add_picture(io.BytesIO(png))
    p = tmp_path / "d.docx"
    d.save(str(p))
    out = tmp_path / "d.json"
    doc, errors = process_single(
        p, output_path=out)
    assert errors == []
    types = [e.type
             for e in doc.elements]
    assert "paragraph" in types
    assert "table" in types
    assert "image" in types
    assert all(
        e.confidence == 0.95
        for e in doc.elements)
    j = json.loads(
        out.read_text(encoding="utf-8"))
    assert all(
        e["confidence"] == 0.95
        for e in j["elements"])
