r"""pipeline DOCX 特殊样式与输出确定性（Round 1585）。

新角度：R1577 锁标题层级——**Title/Quote 家族、
分页符、跨运行字节确定性**零覆盖：

- **Title 样式** → heading；**Subtitle/Quote/
  Intense Quote** → paragraph（样式名保留）
- **分页符** → 仅产生 '(空段落)' 段落元素
- **同一输入两次运行** → 输出 JSON **字节级一致**
  （无时间戳/随机量）
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single
from tests._synthetic_docs import (
    build_minimal_pdf,
)


def test_special_styles_and_break(
        tmp_path):
    d = docxlib.Document()
    d.add_paragraph("T", style="Title")
    d.add_paragraph("S",
                    style="Subtitle")
    d.add_paragraph("Q", style="Quote")
    d.add_paragraph(
        "I", style="Intense Quote")
    d.add_page_break()
    d.add_paragraph("After")
    p = tmp_path / "s.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    got = [(e.type, e.content,
            e.metadata["style"])
           for e in doc.elements]
    assert got == [
        ("heading", "T", "Title"),
        ("paragraph", "S",
         "Subtitle"),
        ("paragraph", "Q", "Quote"),
        ("paragraph", "I",
         "Intense Quote"),
        ("paragraph", "(空段落)",
         "Normal"),
        ("paragraph", "After",
         "Normal")]


def test_output_byte_determinism(
        tmp_path):
    p = build_minimal_pdf(
        tmp_path / "d.pdf",
        text="(DET)")
    o1 = tmp_path / "a.json"
    o2 = tmp_path / "b.json"
    d1, e1 = process_single(
        p, output_path=o1)
    d2, e2 = process_single(
        p, output_path=o2)
    assert e1 == [] and e2 == []
    assert d1.document_id \
        == d2.document_id
    assert o1.read_bytes() \
        == o2.read_bytes()
