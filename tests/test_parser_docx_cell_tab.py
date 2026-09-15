r"""DOCX 表格 cell 内 w:tab——中置 \t 保留、前置 \t 剥除（Round 1954，a 优先级）。

edges49/65 锁 body 段落 w:tab → 字面 '\t' 进 content；R1839
锁 md parser 表格 cell 的 hard tab；**DOCX cell × w:tab** 广
扫（qn("w:tab") × table 各形）实证零覆盖。探针 R1954 实证
（`_Cell.text` 经 `Paragraph.text` tab→'\t'；parser 每 cell
`.strip()` :553）：

- **T1 中置 tab**：'col1\tcol2' → md '| col1\tcol2 | right |'
- **T2 前置 tab**：'\tcol2' → strip 剥首 → '| col2 | right |'
  （**与 body 段落不同——body 不 strip，前置 \t 会留**；
  edges49 的 body 版无此差异面）
- **T3 双 tab**：'a\t\tb' → 内部双 \t 全留

判别式：若 cell 序列化改剥全部 \t 则 T1/T3 翻红；若 strip
取消则 T2 翻红。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _parse(build):
    with tempfile.TemporaryDirectory() as td:
        doc = Document()
        doc.add_paragraph("intro")
        tbl = doc.add_table(rows=1, cols=2)
        build(tbl)
        p = Path(td) / "t.docx"
        doc.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def _add_tabs(cell, n: int):
    r = cell.paragraphs[0].add_run()
    for _ in range(n):
        r.add_tab()


def test_mid_cell_tab_kept():
    """T1：cell 段内中置 w:tab → md 首格字面 '\t'。"""

    def build(tbl):
        c = tbl.cell(0, 0).paragraphs[0]
        c.add_run("col1")
        _add_tabs(tbl.cell(0, 0), 1)
        c.add_run("col2")
        tbl.cell(0, 1).text = "right"

    d = _parse(build)
    tbl_e = d.elements[1]
    assert tbl_e.type == "table"
    assert tbl_e.content == "| col1\tcol2 | right |\n| --- | --- |"
    assert tbl_e.metadata["row_count"] == 1
    assert tbl_e.metadata["col_count"] == 2
    assert d.warnings == []


def test_leading_cell_tab_stripped():
    """T2：cell 以 w:tab 开头 → strip 剥首（body 段落会留）。"""

    def build(tbl):
        _add_tabs(tbl.cell(0, 0), 1)
        tbl.cell(0, 0).paragraphs[0].add_run("col2")
        tbl.cell(0, 1).text = "right"

    d = _parse(build)
    assert d.elements[1].content == "| col2 | right |\n| --- | --- |"
    assert d.warnings == []


def test_double_internal_tabs_kept():
    """T3：连续两个 w:tab → 内部 '\t\t' 全保留。"""

    def build(tbl):
        c = tbl.cell(0, 0).paragraphs[0]
        c.add_run("a")
        _add_tabs(tbl.cell(0, 0), 2)
        c.add_run("b")
        tbl.cell(0, 1).text = "right"

    d = _parse(build)
    assert d.elements[1].content == "| a\t\tb | right |\n| --- | --- |"
    assert d.warnings == []
