r"""DOCX 修订标记（track changes）内容静默缺席（Round 1989，a 优先级）。

真实协作文档常带未接受修订；grep 实证 docx 测试零 w:ins/
w:del 匹配（vanish/webHidden 隐藏文本 R1978 已锁且照提——
修订标记是另一语义：**结构性**排除）。python-docx
Paragraph.text 只拼直接 w:r 子 run（fallback_parser._parse_
docx:468 用 para.text；cell 同理 c.text），w:ins 包裹的
run 与 w:del 里的 w:delText 均非直接子 run。探针 R1989 实证：

- **T1 混排段**：'KEPT ' + w:ins('INSERTED') → 只提 'KEPT'
  ——待接受插入静默丢弃
- **T2 删除标记**：w:del(w:delText 'GONE') + 'REMAINS' →
  只提 'REMAINS'（delText 非 w:t，双保险不出现）
- **T3 段中插入**：'BEFORE' + w:ins('MID') + 'AFTER' →
  'BEFOREAFTER' 无空隙直接拼接（插入处文本断层不可见）
- **T4 表格 cell 内 w:ins** → '| C |'（cell 同路径丢弃）

四态均零告警——修订存在性与隐藏文本一样完全静默。

判别式：若 Paragraph.text 改为遍历全 w:t（含 ins 内）则
T1/T3/T4 文本断言翻红；若 delText 被误当 w:t 消费则 T2 出
现 'GONE' 翻红；若引入修订告警则 warnings 断言翻红。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser

_W = nsdecls("w")


def _ins_wrap(text: str, rid: int) -> str:
    return (f'<w:ins {_W} w:id="{rid}" w:author="probe"'
            f' w:date="2026-01-01T00:00:00Z">'
            f'<w:r><w:t>{text}</w:t></w:r></w:ins>')


def _del_wrap(text: str, rid: int) -> str:
    return (f'<w:del {_W} w:id="{rid}" w:author="probe"'
            f' w:date="2026-01-01T00:00:00Z">'
            f'<w:r><w:delText>{text}</w:delText></w:r></w:del>')


def _parse(path: Path):
    return FallbackParser().parse(path, compute_file_hash(path))


def test_ins_run_dropped(tmp_path):
    """T1：w:ins 包裹的 run 静默丢弃 → 只提 'KEPT'。"""
    p = tmp_path / "t1.docx"
    d0 = Document()
    para = d0.add_paragraph("KEPT ")
    para._p.append(parse_xml(_ins_wrap("INSERTED", 1)))
    d0.save(p)
    d = _parse(p)
    assert [e.content for e in d.elements] == ["KEPT"]
    assert d.elements[0].source_locator["paragraph_index"] == 0
    assert d.warnings == []


def test_del_text_invisible(tmp_path):
    """T2：w:del/w:delText 不可见 → 只提 'REMAINS'。"""
    p = tmp_path / "t2.docx"
    d0 = Document()
    para = d0.add_paragraph("")
    para._p.append(parse_xml(_del_wrap("GONE", 2)))
    para.add_run("REMAINS")
    d0.save(p)
    d = _parse(p)
    assert [e.content for e in d.elements] == ["REMAINS"]
    assert d.warnings == []


def test_mid_paragraph_gapless_concat(tmp_path):
    """T3：段中 ins 缺席 → 'BEFOREAFTER' 无空隙拼接。"""
    p = tmp_path / "t3.docx"
    d0 = Document()
    para = d0.add_paragraph("BEFORE")
    para._p.append(parse_xml(_ins_wrap("MID", 3)))
    para.add_run("AFTER")
    d0.save(p)
    d = _parse(p)
    assert [e.content for e in d.elements] == ["BEFOREAFTER"]
    assert d.warnings == []


def test_cell_ins_dropped(tmp_path):
    """T4：表格 cell 内 w:ins 同路径丢弃 → '| C |'。"""
    p = tmp_path / "t4.docx"
    d0 = Document()
    d0.add_paragraph("lead")
    t = d0.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0].paragraphs[0]
    cell.add_run("C")
    cell._p.append(parse_xml(_ins_wrap("ELLINS", 4)))
    d0.save(p)
    d = _parse(p)
    assert [e.type for e in d.elements] == ["paragraph", "table"]
    assert d.elements[1].content == "| C |\n| --- |"
    assert d.warnings == []
