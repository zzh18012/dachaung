r"""DOCX 空 heading 状态机三边界锁定（Round 1926，a 优先级）。

R1925 锁了"两正文间单空 heading 劈两块"；本轮锁 chunker 硬边界
状态机的三个残余边界（探针 R1926 实证）：

- **连续空 heading** → 3 chunk，中间块**恰为纯占位符单元素块**
  （text=="(空段落)"、source 恰含前一个空 heading）——第二个
  heading 把仅含第一个空 heading 的缓冲封口成块；级数保留
  （空 Heading 2 → level=2）
- **文档开头空 heading**（无前块可封）→ 融合进后续正文**单
  chunk**（R1922 C1 的 PDF 表内 heading 同型，DOCX 侧锁定）
- **文档末尾空 heading** → 收尾 flush 出**纯占位符尾块**
  （单元素 chunk "(空段落)"）

判别式：若 heading 硬边界对空占位 heading 停止触发（或空元素
不进缓冲），三块的拓扑断言翻红；若收尾 flush 丢弃纯占位缓冲，
尾块断言翻红。
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import process_single


def _run(tmp_path: Path, build) -> object:
    p = tmp_path / "s.docx"
    d = build()
    d.save(str(p))
    doc, errors = process_single(p, write_json=False)
    assert doc is not None and errors == []
    return doc


def test_consecutive_empty_headings_pure_placeholder_chunk(tmp_path):
    """正文-空H1-空H2-正文 → 3 chunk：中间块恰为纯占位符单元素块
    （仅前一个空 heading）；级数保留（空 Heading 2 → level=2）。"""
    doc = _run(tmp_path, lambda: _doc([
        ("p", "Before body text."),
        ("h1", ""), ("h2", ""),
        ("p", "After body text."),
    ]))
    assert [e.metadata["level"] for e in doc.elements[1:3]] == [1, 2]
    texts = [c.text for c in doc.chunks]
    assert texts == [
        "Before body text.", "(空段落)", "(空段落) After body text."]
    middle = doc.chunks[1]
    assert middle.source_element_ids == [doc.elements[1].element_id]
    assert doc.chunks[2].source_element_ids == [
        doc.elements[2].element_id, doc.elements[3].element_id]


def test_doc_start_empty_heading_fuses_single_chunk(tmp_path):
    """文档开头空 heading 无前块可封 → 融合进后续正文恰 1 chunk
    （与 R1922 C1 的 PDF 表内 heading 同型）。"""
    doc = _run(tmp_path, lambda: _doc([("h1", ""), ("p", "After body text.")]))
    assert len(doc.chunks) == 1
    assert doc.chunks[0].text == "(空段落) After body text."
    assert doc.chunks[0].source_element_ids == [
        e.element_id for e in doc.elements]


def test_doc_end_empty_heading_trailing_placeholder_chunk(tmp_path):
    """文档末尾空 heading → 收尾 flush 出纯占位符尾块（单元素
    chunk，text 恰 "(空段落)"）。"""
    doc = _run(tmp_path, lambda: _doc([
        ("p", "Only body text."), ("h1", "")]))
    assert [c.text for c in doc.chunks] == ["Only body text.", "(空段落)"]
    assert doc.chunks[1].source_element_ids == [doc.elements[1].element_id]


def _doc(parts: list[tuple[str, str]]) -> object:
    d = docxlib.Document()
    for kind, text in parts:
        style = {"p": None, "h1": "Heading 1", "h2": "Heading 2"}[kind]
        para = d.add_paragraph(text) if style is None \
            else d.add_paragraph(text, style=style)
        assert para.text == text
    return d
