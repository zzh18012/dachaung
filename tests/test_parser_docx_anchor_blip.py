r"""DOCX 浮动图（wp:anchor 带真实 blip）可提取 + hyperlink 内嵌 drawing 可见锁定（Round 1898，a 优先级）。

edges61 锁过"浮动 wp:anchor 不可见"——但其构造的 anchor **无
a:blip 数据**，锁的是空 anchor。探针 R1898 反向判别实证
（`_extract_inline_image_rids` :422-433 用 `drawing.iter(qn("a:blip"))`
后代递归——**不在乎 blip 挂在 wp:inline 还是 wp:anchor 下**，
函数名 "inline" 有误导性）：

- **带真实 blip 的浮动图被提取**：wp:inline 改名换姓成 wp:anchor
  （保留 a:graphic 子树）→ image 元素照常生成（rid、locator 完整）
- **判据是 blip 存在性**：同一文档空 anchor（edges61 形态）+ 数据
  anchor 并存 → 只有数据 anchor 成元素
- **hyperlink 内嵌 drawing 可见**：drawing 的 run 被搬进
  w:hyperlink → `.iter()` 递归照常找到（段落文本不含它，图独立成
  元素）
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _png_1x1() -> bytes:
    def chunk(t: bytes, d: bytes) -> bytes:
        return (struct.pack(">I", len(d)) + t + d
                + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF))
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\xff\x00\x00"))
            + chunk(b"IEND", b""))


def _anchorize(paragraph) -> str:
    """段内首个 wp:inline 换成 wp:anchor（保留 a:graphic 子树），返回 rid。"""
    drawing = paragraph._p.find(f".//{qn('w:drawing')}")
    inline = drawing.find(qn("wp:inline"))
    rid = drawing.find(f".//{qn('a:blip')}").get(qn("r:embed"))
    anchor = OxmlElement("wp:anchor")
    for child in list(inline):
        anchor.append(child)
    drawing.replace(inline, anchor)
    return rid


def test_docx_floating_anchor_with_real_blip_is_extracted(tmp_path):
    """判别式：wp:anchor 带真实 blip → image 元素照常（edges61 只锁过
    空 anchor 不可见）；段落文本与零告警不受影响。"""
    doc = Document()
    p = doc.add_paragraph("floating with data")
    p.add_run().add_picture(io.BytesIO(_png_1x1()))
    rid = _anchorize(p)
    path = tmp_path / "anchor.docx"
    doc.save(path)
    parsed = FallbackParser().parse(path, compute_file_hash(path))
    images = [e for e in parsed.elements if e.type == "image"]
    assert len(images) == 1
    assert images[0].source_locator["relationship_id"] == rid
    assert images[0].resource_path == "(unsaved)"
    assert [e.content for e in parsed.elements if e.type == "paragraph"] == [
        "floating with data"]
    assert [w.code for w in parsed.warnings] == []


def test_docx_blip_presence_is_the_only_criterion(tmp_path):
    """同文档并存：空 anchor（无 blip）+ 数据 anchor → 恰 1 个 image
    元素且来自后者——inline/anchor 形态无关，blip 存在性是唯一判据。"""
    doc = Document()
    p_empty = doc.add_paragraph("empty anchor para")
    r = p_empty.add_run()
    drawing = OxmlElement("w:drawing")
    drawing.append(OxmlElement("wp:anchor"))
    r._r.append(drawing)
    p_data = doc.add_paragraph("data anchor para")
    p_data.add_run().add_picture(io.BytesIO(_png_1x1()))
    _anchorize(p_data)
    path = tmp_path / "mixed.docx"
    doc.save(path)
    parsed = FallbackParser().parse(path, compute_file_hash(path))
    images = [e for e in parsed.elements if e.type == "image"]
    assert len(images) == 1
    assert images[0].source_locator["paragraph_index"] == 1  # 数据 anchor 段
    assert [w.code for w in parsed.warnings] == []


def test_docx_drawing_inside_hyperlink_is_extracted(tmp_path):
    """drawing 的 run 搬进 w:hyperlink → .iter() 递归照常找到；段落
    文本不含图、图独立成元素。"""
    doc = Document()
    p = doc.add_paragraph("pic in hyperlink")
    run = p.add_run()
    run.add_picture(io.BytesIO(_png_1x1()))
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), "rId9")
    p._p.replace(run._r, hl)
    hl.append(run._r)
    path = tmp_path / "hlimg.docx"
    doc.save(path)
    parsed = FallbackParser().parse(path, compute_file_hash(path))
    images = [e for e in parsed.elements if e.type == "image"]
    assert len(images) == 1
    assert images[0].source_locator["relationship_id"] == "rId9"
    assert [e.content for e in parsed.elements if e.type == "paragraph"] == [
        "pic in hyperlink"]
    assert [w.code for w in parsed.warnings] == []
