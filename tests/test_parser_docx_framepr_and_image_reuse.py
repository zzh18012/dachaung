r"""DOCX framePr 定位段 + 同 rId 图复用（Round 1970，a 优先级）。

framePr（旧式定位文本框段落）零覆盖（广扫零匹配）；
R1969 锁了 PDF 对象别名，DOCX 侧同图 rId 两处引用语义
未知。探针 R1970 实证：

- **F1 w:framePr 段**（x/y/w/h 页锚定）→ 文本照提为普通
  paragraph、pidx 0、定位**全忽略**、零告警
- **F2 同一 rId 的 w:drawing 深拷贝进两段** → **不去重**：
  两个 image 元素 + 两个独立 PNG（前缀 para0/para1 + 全
  局 counter 00/01）——与 PDF A2 页共享对偶
- **F3 framePr 段夹在普通段之间** → 文档序（非视觉序）
  0/1/2 顺延

判别式：若按 rId 去重则 F2 单元素单文件翻红；若实现框
定位重排则 F3 顺序变翻红。
"""

from __future__ import annotations

import copy
import io
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser

_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
    "53de0000000c4944415408d763f8ffff3f0005fe02fea735c9d40000000049"
    "454e44ae426082")

_FRAME_INNER = ('<w:framePr w:w="2000" w:h="1000" w:x="1000" '
                'w:y="2000" w:hAnchor="page" w:vAnchor="page"/>')


def _framed(d: Document, text: str):
    p = d.add_paragraph()
    p._p.append(parse_xml(f"<w:pPr {nsdecls('w')}>{_FRAME_INNER}</w:pPr>"))
    p.add_run(text)
    return p


def _parse(kind: str):
    d = Document()
    outdir = None
    if kind == "single":
        _framed(d, "FRAMED TEXT")
    elif kind == "reuse":
        p1 = d.add_paragraph()
        p1.add_run("first ")
        p1.add_run().add_picture(io.BytesIO(_PNG))
        p2 = d.add_paragraph()
        p2.add_run("second ")
        drawing = p1._p.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")
        r2 = parse_xml(f'<w:r {nsdecls("w")}></w:r>')
        r2.append(copy.deepcopy(drawing[0]))
        p2._p.append(r2)
    else:  # order
        d.add_paragraph("before frame")
        _framed(d, "MIDDLE FRAME")
        d.add_paragraph("after frame")
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "f.docx"
        d.save(p)
        outdir = Path(td) / "img" if kind == "reuse" else None
        parser = (FallbackParser(image_output_dir=outdir) if outdir
                  else FallbackParser())
        doc = parser.parse(p, compute_file_hash(p))
        files = sorted(f.name for f in outdir.iterdir()) if outdir else []
        return doc, files


def test_framepr_paragraph_visible_position_ignored():
    """F1：framePr 段 → 普通 paragraph 'FRAMED TEXT'、pidx
    0、定位全忽略、零告警。"""
    d, _ = _parse("single")
    assert [e.type for e in d.elements] == ["paragraph"]
    assert d.elements[0].content == "FRAMED TEXT"
    assert d.elements[0].source_locator["paragraph_index"] == 0
    assert d.warnings == []


def test_same_rid_two_paragraphs_two_elements():
    """F2：同 rId 深拷贝进两段 → 不去重：两 image 元素、两
    独立 PNG（para0_00 / para1_01，全局 counter）。"""
    d, files = _parse("reuse")
    assert [e.type for e in d.elements] == [
        "paragraph", "image", "paragraph", "image"]
    imgs = [e for e in d.elements if e.type == "image"]
    assert [e.source_locator["paragraph_index"] for e in imgs] == [0, 1]
    assert all(e.metadata["extracted_to_disk"] for e in imgs)
    assert files == [
        f for f in files if f.endswith(("_para0_00.png", "_para1_01.png"))]
    assert len(files) == 2
    assert d.warnings == []


def test_framepr_between_paragraphs_document_order():
    """F3：framePr 夹中间 → 文档序 0/1/2（非视觉定位序）。"""
    d, _ = _parse("order")
    assert [e.content for e in d.elements] == [
        "before frame", "MIDDLE FRAME", "after frame"]
    assert [e.source_locator["paragraph_index"] for e in d.elements] == [0, 1, 2]
    assert d.warnings == []
