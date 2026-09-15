r"""DOCX 表格 cell 内 hyperlink——可见文本并入、悬挂 r:id 亦可（Round 1957，a 优先级）。

edges15 锁 body 段落 hyperlink 可见文本并入（悬挂 r:id 无
relationship 亦可解析）；**cell × hyperlink** 经广扫
（add_hyperlink/hyperlink×cell/表格×链接 各形）实证零覆盖。
探针 R1957 实证（`_Cell.text` 经 `Paragraph.text`，python-docx
1.2.0 含 hyperlink 可见文本；parser 只读 c.text）：

- **H1 cell 中置**：run 'see ' + link 'docs here' → md
  '| see docs here | right |'、无 link 元数据、零告警
- **H2 链后接 run**：'see ' + link 'docs' + ' now' →
  '| see docs now | right |'（序保真）
- **H3 链独占 cell**：cell 文本全部来自 hyperlink →
  '| docs here | right |'（无 run 也不空）

判别式：若 cell 通道改只读 w:r 直系 run（跳 hyperlink）则
三断言齐翻红（H3 空 cell）。
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser


def _add_link(par, text: str) -> None:
    hl = par._p.makeelement(qn("w:hyperlink"),
                            {qn("r:id"): "rId9"})
    r = hl.makeelement(qn("w:r"), {})
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    r.append(t)
    hl.append(r)
    par._p.append(hl)


def _parse(build):
    with tempfile.TemporaryDirectory() as td:
        d = Document()
        d.add_paragraph("intro")
        tbl = d.add_table(rows=1, cols=2)
        build(tbl.cell(0, 0).paragraphs[0])
        tbl.cell(0, 1).text = "right"
        p = Path(td) / "t.docx"
        d.save(p)
        return FallbackParser().parse(p, compute_file_hash(p))


def test_mid_cell_hyperlink_text_merged():
    """H1：run + 悬挂 r:id hyperlink → '| see docs here |'。"""
    d = _parse(lambda par: (
        par.add_run("see "), _add_link(par, "docs here")))
    tbl_e = d.elements[1]
    assert tbl_e.type == "table"
    assert tbl_e.content == "| see docs here | right |\n| --- | --- |"
    assert tbl_e.metadata["row_count"] == 1
    assert tbl_e.metadata["col_count"] == 2
    assert d.warnings == []


def test_run_after_hyperlink_order_kept():
    """H2：link 后再接 run → '| see docs now |'（序保真）。"""
    d = _parse(lambda par: (
        par.add_run("see "), _add_link(par, "docs"),
        par.add_run(" now")))
    assert d.elements[1].content == "| see docs now | right |\n| --- | --- |"
    assert d.warnings == []


def test_hyperlink_only_cell():
    """H3：cell 文本全部来自 hyperlink → '| docs here |'。"""
    d = _parse(lambda par: _add_link(par, "docs here"))
    assert d.elements[1].content == "| docs here | right |\n| --- | --- |"
    assert d.warnings == []
