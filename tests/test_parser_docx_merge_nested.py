r"""DOCX 合并单元格重复展开 + 嵌套表丢弃锁定（Round 1895，a 优先级）。

grep 零覆盖（parser 测试无 gridSpan/vMerge/嵌套表），现实极常见。
探针 R1895 三组实证（python-docx `row.cells` 语义 × `_parse_docx`
:545-570 平铺不递归）：

- **横向合并按跨度重复**：cell(0,0).merge(cell(0,1)) 后
  `row.cells` 对每个网格列返回同一 cell → "TL\\nTR" 在 markdown
  行里出现两次（且合并文本自带 \\n，破坏 markdown 一行一行的
  表结构）
- **纵向合并跨行重复**：vMerge 续行 `row.cells` 返回原 cell →
  合并文本 "TOP\\nBOT" 在两行 markdown 里都出现
- **嵌套表整体丢弃**：w:tc 里 add_table 的内表既不进 `c.text`
  也不进 body.iterchildren（:463 只走直接子节点）→ "inner cell"
  无任何元素承载（R1889 sdt 丢弃的同族，XML 形态不同）
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from app.parsers.fallback_parser import _parse_docx


def _hmerge_docx(path: Path) -> None:
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "TL"
    tbl.cell(0, 1).text = "TR"
    tbl.cell(1, 0).text = "BL"
    tbl.cell(1, 1).text = "BR"
    tbl.cell(0, 0).merge(tbl.cell(0, 1))
    doc.save(path)


def test_docx_horizontal_merge_duplicates_cell_across_columns(tmp_path):
    """横向合并：row.cells 每个网格列给同一 cell → 合并文本按跨度
    重复（'| TL\\nTR | TL\\nTR |'），合并自带的 \\n 留在 md 里。"""
    path = tmp_path / "hmerge.docx"
    _hmerge_docx(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    assert len(elements) == 1
    assert elements[0].type == "table"
    assert elements[0].content == "| TL\nTR | TL\nTR |\n| --- | --- |\n| BL | BR |"
    # 重复展开不影响网格计数：仍是 2 列
    assert elements[0].metadata["col_count"] == 2


def test_docx_vertical_merge_repeats_text_in_both_rows(tmp_path):
    """纵向合并：vMerge 续行返回原 cell → 合并文本 "TOP\\nBOT" 在
    两个 markdown 行都出现（2×2 合一列后信息不丢但翻倍）。"""
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "TOP"
    tbl.cell(0, 1).text = "TR"
    tbl.cell(1, 0).text = "BOT"
    tbl.cell(1, 1).text = "BR"
    tbl.cell(0, 0).merge(tbl.cell(1, 0))
    path = tmp_path / "vmerge.docx"
    doc.save(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    md = elements[0].content
    assert md == "| TOP\nBOT | TR |\n| --- | --- |\n| TOP\nBOT | BR |"
    assert elements[0].metadata["row_count"] == 2


def test_docx_nested_table_content_dropped(tmp_path):
    """嵌套表：cell 内 add_table 的内容既不进 c.text 也不成独立
    元素——"inner cell" 无处承载，md 只有 outer 文本。"""
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    tbl.cell(0, 0).text = "outer cell"
    inner = tbl.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "inner cell"
    path = tmp_path / "nested.docx"
    doc.save(path)
    elements, warnings = _parse_docx(path, "sha" * 21, "doc-x", None)
    assert warnings == []
    assert len(elements) == 1
    assert elements[0].content == "| outer cell |\n| --- |"
    assert "inner cell" not in elements[0].content
    assert elements[0].metadata["row_count"] == 1
    assert elements[0].metadata["col_count"] == 1
