r"""pipeline Unicode 内容：WinAnsi 八进制与 CJK（Round 1586）。

新角度：R1585 锁确定性——**多字节内容编码与
无空白长文切分**零覆盖：

- **WinAnsi 八进制转义**（\\351 \\357 \\374）→
  正确解码 'café naïve über'
- **DOCX 中文** → 元素与 JSON 往返无损、校验通过
- **纯 CJK 850 字无空白** → 无词边界可用、
  **硬切** [800, 50]（strategy 仍 sentence_split）
"""

from __future__ import annotations

from pathlib import Path

import docx as docxlib

from app.pipeline import (
    process_single,
    validate_only,
)

_FONT = ("<< /Type /Font"
         " /Subtype /Type1"
         " /BaseFont /Helvetica"
         " /Encoding"
         " /WinAnsiEncoding >>")


def _pdf(tmp_path: Path, name: str,
         c: str) -> Path:
    objs = {
        1: "<< /Type /Catalog"
           " /Pages 2 0 R >>",
        2: "<< /Type /Pages"
           " /Kids [3 0 R]"
           " /Count 1 >>",
        3: ("<< /Type /Page"
            " /Parent 2 0 R"
            " /MediaBox"
            " [0 0 612 792]"
            " /Resources << /Font"
            " << /F1 5 0 R >> >>"
            " /Contents 4 0 R >>"),
        4: f"<< /Length {len(c)} >>"
           f"\nstream\n{c}\nendstream",
        5: _FONT,
    }
    pdf = b"%PDF-1.4\n"
    for oid in sorted(objs):
        o = objs[oid]
        if isinstance(o, str):
            o = o.encode("latin-1")
        pdf += (f"{oid} 0 obj\n"
                ).encode() + o \
            + b"\nendobj\n"
    pdf += (b"trailer"
            b" << /Root 1 0 R"
            b" /Size 6 >>\n%%EOF")
    p = tmp_path / name
    p.write_bytes(pdf)
    return p


def test_winansi_octal_decoded(
        tmp_path):
    c = ("BT /F1 12 Tf 72 700"
         " Td (caf\\351"
         " na\\357ve"
         " \\374ber) Tj ET")
    p = _pdf(tmp_path, "u.pdf", c)
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert el.content == (
        "café naïve über")
    (ch,) = doc.chunks
    assert ch.text == "café naïve über"


def test_docx_cjk_roundtrip(
        tmp_path):
    d = docxlib.Document()
    d.add_paragraph("中文内容测试段落")
    d.add_heading("标题中文", level=1)
    p = tmp_path / "c.docx"
    d.save(str(p))
    out = tmp_path / "c.json"
    doc, errors = process_single(
        p, output_path=out)
    assert errors == []
    assert [(e.type, e.content)
            for e in doc.elements] == [
        ("paragraph", "中文内容测试段落"),
        ("heading", "标题中文")]
    ok, msg = validate_only(out)
    assert ok, msg


def test_cjk_hard_split(
        tmp_path):
    d = docxlib.Document()
    d.add_paragraph("汉" * 850)
    p = tmp_path / "l.docx"
    d.save(str(p))
    doc, errors = process_single(
        p, write_json=False)
    assert errors == []
    (el,) = doc.elements
    assert len(el.content) == 850
    got = [(len(c.text),
            c.metadata["strategy"])
           for c in doc.chunks]
    assert got == [
        (800,
         "long_paragraph_sentence_split"),
        (50,
         "long_paragraph_sentence_split")]
    assert doc.chunks[1].text \
        == "汉" * 50
