r"""DOCX VML 图片与 OLE 对象不可见（Round 1990，a 优先级）。

老版 Word 文档常见 VML 行内图（w:pict > v:shape >
v:imagedata r:id）与 OLE 嵌入（w:object 包 v:shape +
o:OLEObject，如 Equation.3 / Excel 片段）。grep 实证
imagedata/OLE/pict 在 parser 与全部测试零匹配（图片既有
覆盖全走 w:drawing：R1946 格式透传、行内图融合、浮动图
丢弃）。fallback 只扫 w:drawing（fallback_parser:427
paragraph.iter(qn("w:drawing"))）。探针 R1990 实证（rId
与 media part 真实存在于包内，仅换引用形态）：

- **T1 VML-only 文档**：drawing run 剥除、同 rId 改挂
  v:imagedata → 零 image 元素、单 '(空段落)' 占位——图片
  字节在包里也照样丢
- **T2 w:object OLE**：前导文本照提、零 image 元素——
  OLEObject/v:imagedata 双不可见
- **T3 对照 w:drawing**（R1974 规则：证分支真实触发）：
  同一图片 add_picture → image 元素照发（ext 'png'、
  byte_size 逐字节）

三态均零告警。

判别式：若 parser 扩展扫 v:imagedata/w:object 则 T1/T2
出现 image 元素翻红；若 rels 存在性参与判定（无 drawing 但
有关系）则 T1 翻红；T3 翻红说明夹具坏（图片路径本身断）。
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from app.hash import compute_file_hash
from app.parsers.fallback_parser import FallbackParser

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
       b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
       b"\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")

_W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
      ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006'
      '/relationships"'
      ' xmlns:v="urn:schemas-microsoft-com:vml"'
      ' xmlns:o="urn:schemas-microsoft-com:office:office"')


def _pict_xml(rid: str) -> str:
    return (f'<w:pict {_W}><v:shape id="vml1"'
            f' style="width:10pt;height:10pt">'
            f'<v:imagedata r:id="{rid}"/></v:shape></w:pict>')


def _object_xml(rid: str) -> str:
    return (f'<w:object {_W}><v:shape id="vml2"'
            f' style="width:10pt;height:10pt">'
            f'<v:imagedata r:id="{rid}"/></v:shape>'
            f'<o:OLEObject Type="Embed" ProgID="Equation.3"'
            f' ShapeID="vml2" DrawAspect="Icon" r:id="{rid}"/>'
            f'</w:object>')


def _strip_drawing_runs(d: Document) -> str:
    rid = d.inline_shapes[-1]._inline.graphic.graphicData.pic \
        .blipFill.blip.embed
    p0 = d.paragraphs[0]._p
    for r in list(p0.findall(qn("w:r"))):
        p0.remove(r)
    return rid


def test_vml_imagedata_invisible(tmp_path):
    """T1：v:imagedata（rId/media 真实存在）→ 零 image 元素。"""
    p = tmp_path / "vml.docx"
    d0 = Document()
    d0.add_picture(io.BytesIO(PNG))
    rid = _strip_drawing_runs(d0)
    d0.paragraphs[0]._p.append(parse_xml(_pict_xml(rid)))
    d0.save(p)
    d = FallbackParser().parse(p, compute_file_hash(p))
    assert [(e.type, e.content) for e in d.elements] == [
        ("paragraph", "(空段落)")]
    assert all(e.type != "image" for e in d.elements)
    assert d.warnings == []


def test_ole_object_text_only(tmp_path):
    """T2：w:object OLE → 文本照提、零 image 元素。"""
    p = tmp_path / "ole.docx"
    d0 = Document()
    d0.add_picture(io.BytesIO(PNG))
    rid = _strip_drawing_runs(d0)
    d0.paragraphs[0].add_run("OLE HERE ")
    d0.paragraphs[0]._p.append(parse_xml(_object_xml(rid)))
    d0.save(p)
    d = FallbackParser().parse(p, compute_file_hash(p))
    assert [(e.type, e.content) for e in d.elements] == [
        ("paragraph", "OLE HERE")]
    assert all(e.type != "image" for e in d.elements)
    assert d.warnings == []


def test_drawing_control_fires(tmp_path):
    """T3：对照 w:drawing 同一图片 → image 元素照发（分支活）。"""
    p = tmp_path / "ctl.docx"
    d0 = Document()
    d0.add_picture(io.BytesIO(PNG))
    d0.save(p)
    d = FallbackParser().parse(p, compute_file_hash(p))
    imgs = [e for e in d.elements if e.type == "image"]
    assert len(imgs) == 1
    assert imgs[0].metadata["ext"] == "png"
    assert imgs[0].metadata["byte_size"] == len(PNG)
    assert d.warnings == []
